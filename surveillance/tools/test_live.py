# -*- coding: utf-8 -*-
"""第2弾：入力を書き換えたときに判定が追従するか（＝『入力するたびに反映』）を検証"""
import datetime as dt
import openpyxl, subprocess, sys
from common import *

D = lambda m, d: dt.datetime(2025, m, d)
ok = ng = 0


def chk(label, got, want):
    global ok, ng
    g = got.date() if isinstance(got, dt.datetime) else got
    w = want.date() if isinstance(want, dt.datetime) else want
    if g == w or (g in (None, "") and w in (None, "")):
        ok += 1; print(f"  ok   {label}: {got!r}")
    else:
        ng += 1; print(f"  NG   {label}: got {got!r} want {want!r}")


def mutate(src, dst):
    wb = openpyxl.load_workbook(src)
    al, cl, lst = wb["All"], wb["臨床所見"], wb["VAE対象一覧"]

    # ① ブロック3：FiO2の上昇幅を 20→5 ポイントに下げる → VAC は消えるはず
    for d in (4, 5, 6):
        al.cell(all_f(3), C0 + d - 1, 45)

    # ② ブロック1：PEEP上昇を 3→2 cmH2O に下げる → VAC は消えるはず
    for d in (5, 6):
        al.cell(all_p(1), C0 + d - 1, 7)

    # ③ ブロック6：第2エピソード（4/8開始・8日間）に FiO2悪化を入れる
    #    MV日 1,2=40 / 3,4=70 → 第2エピソードのMV3日目がDOE = 4/10
    for d, v in ((8, 40), (9, 40), (10, 70), (11, 70), (12, 40),
                 (13, 40), (14, 40), (15, 40)):
        al.cell(all_f(6), C0 + d - 1, v)
    wb["VAE-06"]["C9"] = 2                      # 対象エピソードを2件目に切替

    # ④ ブロック9：PVAP基準を消す → PVAP から IVAC に格下げされるはず
    cl.cell(clin_base(9) + 6, C0 + 5 - 1).value = None

    # ⑤ 新規患者をブロック20に追加 → VAE-12 に自動で割り当たるはず
    al.cell(all_id(20), 1, "T20")
    al.cell(all_name(20), 1, "検証20")
    #    FiO2 40,40,40,70,70 → ベースラインはMV2-3日目、DOEはMV4日目（4/4）
    for d, (f, p) in enumerate([(40, 5), (40, 5), (40, 5), (70, 5), (70, 5)]):
        al.cell(all_v(20), C0 + d, "V有")
        al.cell(all_f(20), C0 + d, f)
        al.cell(all_p(20), C0 + d, p)
    lst.cell(LIST_B_TOP + 20 - 1, 8, 65)

    wb.save(dst)
    print("wrote", dst)


def verify(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    lst = wb["VAE対象一覧"]

    print("\n① ブロック3：FiO2の上昇を20→5ポイントに変更")
    chk("  判定①が消える", wb["VAE-03"]["K9"].value, "")

    print("\n② ブロック1：PEEPの上昇を3→2cmH2Oに変更")
    chk("  判定①が消える", wb["VAE-01"]["K9"].value, "")

    print("\n③ ブロック6：第2エピソード（C9=2）に切替")
    w = wb["VAE-06"]
    chk("  MV1日目", w["C10"].value, D(4, 8))
    chk("  本エピソードMV日数", w["K5"].value, 8)
    chk("  DOE①", w["K8"].value, D(4, 10))
    chk("  判定①", w["K9"].value, "VAC")

    print("\n④ ブロック9：PVAP基準を削除 → IVACへ格下げ")
    chk("  判定①", wb["VAE-08"]["K9"].value, "IVAC")

    print("\n⑤ ブロック20に患者を追加 → VAE-12 が自動で埋まる")
    w = wb["VAE-12"]
    chk("  ブロック番号", w["C4"].value, 20)
    chk("  患者ID", w["C5"].value, "T20")
    chk("  MV1日目", w["C10"].value, D(4, 1))
    chk("  DOE①", w["K8"].value, D(4, 4))
    chk("  判定①", w["K9"].value, "VAC")

    print("\n⑥ 一覧セクションB：初回／最終MV日")
    r5 = LIST_B_TOP + 5 - 1
    chk("  ブロック5 初回MV日", lst.cell(r5, 5).value, D(4, 1))
    chk("  ブロック5 最終MV日", lst.cell(r5, 6).value, D(4, 20))
    r6 = LIST_B_TOP + 6 - 1
    chk("  ブロック6 最終MV日", lst.cell(r6, 6).value, D(4, 15))

    print(f"\n==== ok={ok}  NG={ng} ====")
    return ng


if __name__ == "__main__":
    if sys.argv[1] == "mutate":
        mutate(sys.argv[2], sys.argv[3])
    else:
        sys.exit(1 if verify(sys.argv[2]) else 0)
