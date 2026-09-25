# -*- coding: utf-8 -*-
"""臨床所見シート

VAE対象一覧セクションAで VAC と判定された患者だけにスロットが対応する。
スロット n ＝ VAC患者の n 人目（DOEの早い順）。ここに入力した患者だけが個別判定シートを持つ。

入力しやすくするための見た目：
  ・項目のまとまり（体温／WBC／抗菌薬／PVAP）ごとに色分けし、スロットの境目を太線で囲む
  ・DOE±2日（VAEウィンドウ＝入力する日）の列を黄色、DOE当日を橙色で示す（条件付き書式）
  ・月の変わり目に縦線、4行目に「yyyy年m月」
  ・A列の「▶DOE前後へ」で、365日分の横スクロールをせずにDOEの列へ移動できる
"""
import datetime as dt
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter as gl
from common import *

# A列（スロットの見出し）の各行
META_SLOT, META_BLOCK, META_ID, META_NAME, META_DOE1, META_WARN, META_DOE2, META_LINK = range(8)

LINE_SLOT = Side(style="medium", color=NAVY)       # スロットの境目
LINE_GROUP = Side(style="thin", color="8EA9C1")     # 項目のまとまりの境目
LINE_ROW = Side(style="hair", color="BFBFBF")       # まとまりの中の行
LINE_DAY = Side(style="thin", color="D9D9D9")       # 日の境目
LINE_MONTH = Side(style="medium", color="808080")   # 月の変わり目


def _group(off):
    for g0, n, name, strong, pale in CLIN_GROUPS:
        if g0 <= off < g0 + n:
            return g0, name, strong, pale
    raise ValueError(off)


def _top(off):
    if off == 0:
        return LINE_SLOT
    return LINE_GROUP if _group(off)[0] == off else LINE_ROW


def month_starts(ws_all):
    """日付列のうち、月の初日の列（年度初日を含む）"""
    out = []
    prev = None
    for c in range(C0, CN + 1):
        v = ws_all.cell(DATE_ROW, c).value
        m = (v.year, v.month) if isinstance(v, (dt.date, dt.datetime)) else None
        if m != prev:
            out.append(c)
        prev = m
    return out


def build(wb):
    ws = wb.create_sheet("臨床所見")
    ws.sheet_properties.tabColor = "C55A11"
    ws.sheet_view.showGridLines = False
    LAST = CLIN_LAST_ROW
    mstart = month_starts(wb["All"])
    mset = set(mstart)

    ws["A1"] = "臨床所見（IVAC／PVAP 判定用）　― VACと判定された患者だけを入力します ―"
    ws["A1"].font = F_TITLE
    ws["A1"].fill = fill("C55A11")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=40)
    ws.row_dimensions[1].height = 27.75

    # ---- 凡例（A・B列は固定表示なので横にスクロールしても見える）-------------
    legend = [(2, 1, "体温", CLIN_GROUPS[0][3]), (2, 2, "WBC", CLIN_GROUPS[1][3]),
              (3, 1, "抗菌薬", CLIN_GROUPS[2][3]), (3, 2, "PVAP基準", CLIN_GROUPS[3][3]),
              (4, 1, "DOE±2日＝入力", WINDOW), (4, 2, "DOE当日", WINDOW_DOE)]
    for r, c, text, color in legend:
        style_cell(ws.cell(r, c, text), color, Font(bold=True, size=9), A_C)

    notes = [
        "黄色の列（DOE±2日＝VAEウィンドウ）に入力します。体温・WBCはその日の最高値・最低値、"
        "新規抗菌薬は開始した日に「はい」、PVAP基準は該当した日に選びます。橙色がDOE当日です。",
        "スロット番号＝VAE対象一覧セクションAのNo.（VACと判定された順＝DOEの早い順）。"
        "A列の「▶DOE前後へ」をクリックするとDOEの列へ移動します。灰色のスロットは未割当で、入力は不要です。"
        "新規抗菌薬＝NHSN適格抗菌薬を開始した日（開始日の前2日に同薬の投与がなく、4QAD以上継続）。",
    ]
    for i, text in enumerate(notes):
        c = ws.cell(2 + i, 3, text)
        c.font = F_NOTE
        c.alignment = A_L
        ws.merge_cells(start_row=2 + i, start_column=3, end_row=2 + i, end_column=80)

    # ---- 4行目：月、5行目：日付 ---------------------------------------
    ws["A5"] = "患者"
    ws["B5"] = "項目"
    for a in ("A5", "B5"):
        style_cell(ws[a], NAVY, Font(bold=True, size=9, color="FFFFFF"), A_CW)
    bounds = mstart + [CN + 1]
    for i in range(len(mstart)):
        c0, c1 = bounds[i], bounds[i + 1] - 1
        cell = ws.cell(4, c0, f'=TEXT(All!{gl(c0)}${DATE_ROW},"yyyy年m月")')
        style_cell(cell, BAND, Font(bold=True, size=9, color=NAVY), A_L)
        cell.border = Border(left=LINE_MONTH, top=thin, bottom=thin)
        if c1 > c0:
            ws.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c1)
    for c in range(C0, CN + 1):
        cell = ws.cell(5, c)
        cell.value = f"=All!{gl(c)}{DATE_ROW}"
        style_cell(cell, NAVY, Font(bold=True, size=8, color="FFFFFF"), A_C, numfmt="m/d")
        if c in mset:
            cell.border = Border(left=LINE_MONTH)
    ws.row_dimensions[5].height = 24

    # ---- 入力欄の書式（色分けと罫線）。組合せごとにスタイルを作り回して軽くする ----
    styles = {}

    def input_style(off, month, last_row):
        key = (off, month, last_row)
        if key not in styles:
            _, _, _, pale = _group(off)
            styles[key] = (
                fill(pale),
                Border(left=LINE_MONTH if month else LINE_DAY, right=LINE_DAY,
                       top=_top(off), bottom=LINE_SLOT if last_row else None),
            )
        return styles[key]

    # ---- スロット（VAC患者 n 人目）------------------------------------
    L = "'VAE対象一覧'"
    dates = f"$C$5:${gl(CN)}$5"
    for n in range(1, NCLIN + 1):
        b = clin_base(n)
        ar = LIST_A_TOP + n - 1                  # 一覧セクションAの対応行
        last_slot = (n == NCLIN)
        meta = {
            META_SLOT: (n, '"スロット"0'),
            META_BLOCK: (f'=IF({L}!$B${ar}="","",{L}!$B${ar})', '"ブロック"0'),
            META_ID: (f'=IF({L}!$C${ar}="","",{L}!$C${ar})', "General"),
            META_NAME: (f'=IF({L}!$D${ar}="","",{L}!$D${ar})', "General"),
            META_DOE1: (f'=IF({L}!$H${ar}="","",{L}!$H${ar})', '"DOE① "yyyy/m/d'),
            META_DOE2: (f'=IF({L}!$J${ar}="","",{L}!$J${ar})', '"DOE② "yyyy/m/d'),
        }
        for off, (v, fmt) in meta.items():
            c = ws.cell(b + off, 1, v)
            style_cell(c, AUTO, Font(size=9), A_C, numfmt=fmt)
        style_cell(ws.cell(b + META_SLOT, 1), NAVY, Font(bold=True, size=10, color="FFFFFF"), A_C)

        # 入力済みなのにDOE前後（①・②とも）に値がない＝患者の割当がずれた可能性がある
        def window_count(doe_cell):
            # COUNTA はエラー値そのものを「1件」と数えるため、IFERRORで包むのではなく
            # DOEが日付列に見つかるときだけ数える（DOE②が空のスロットで警告が消えないように）
            m = f"MATCH({doe_cell},All!$C${DATE_ROW}:${gl(CN)}${DATE_ROW},0)"
            d = f"{m}+{C0-1}"
            return (f'IF(ISNUMBER({m}),COUNTA(INDEX({N_CLIN},{b},MAX({C0},{d}-2)):'
                    f'INDEX({N_CLIN},{b+7},MIN({CN},{d}+2))),0)')
        ws.cell(b + META_WARN, 1, (
            f'=IF(COUNTA($C{b}:${gl(CN)}{b+7})=0,"",'
            f'IF($A${b+META_DOE1}="","⚠DOE不明",'
            f'IF({window_count(f"$A${b+META_DOE1}")}+{window_count(f"$A${b+META_DOE2}")}>0,"",'
            f'"⚠割当確認")))'))
        style_cell(ws.cell(b + META_WARN, 1), AUTO, Font(size=8, bold=True, color=RED), A_C)

        # DOEの2日前の列へ移動するリンク（365列を横スクロールしなくて済むように）
        ws.cell(b + META_LINK, 1, (
            f'=IF($A${b+META_DOE1}="","",HYPERLINK("#\'臨床所見\'!"&ADDRESS({b},'
            f'MAX({C0},MATCH($A${b+META_DOE1},{dates},0)+{C0-1}-2)),"▶DOE前後へ"))'))
        style_cell(ws.cell(b + META_LINK, 1), AUTO,
                   Font(size=9, color="0563C1", underline="single"), A_C)

        for off, (label, _kind) in enumerate(CLIN_ROWS):
            g0, _name, strong, _pale = _group(off)
            cell = ws.cell(b + off, 2, label)
            style_cell(cell, strong, Font(size=9, bold=(off == g0)), A_L)
            last_row = last_slot and off == 7
            for col in (1, 2):
                c = ws.cell(b + off, col)
                c.border = Border(left=thin, right=thin if col == 1 else LINE_MONTH,
                                  top=_top(off), bottom=LINE_SLOT if last_row else None)
            for c in range(C0, CN + 1):
                f_, bd = input_style(off, c in mset, last_row)
                cell = ws.cell(b + off, c)
                cell.fill = f_
                cell.border = bd

    # ---- 条件付き書式（上から順に優先）--------------------------------
    rng = f"C6:{gl(CN)}{LAST}"
    base = f"8*INT((ROW()-6)/8)"
    blk = f"INDEX($A$6:$A${LAST},{base}+{META_BLOCK+1})"
    d1 = f"INDEX($A$6:$A${LAST},{base}+{META_DOE1+1})"
    d2 = f"INDEX($A$6:$A${LAST},{base}+{META_DOE2+1})"
    # 1) 患者が割り当てられていないスロットは灰色（入力対象外）
    ws.conditional_formatting.add(
        f"A6:{gl(CN)}{LAST}",
        Rule(type="expression", formula=[f'{blk}=""'],
             dxf=DifferentialStyle(fill=cf_fill("E7E6E6"), font=Font(color="A6A6A6")),
             stopIfTrue=True))
    # 2) DOE当日（①・②）
    ws.conditional_formatting.add(
        rng,
        Rule(type="expression",
             formula=[f'OR(AND({d1}<>"",C$5={d1}),AND({d2}<>"",C$5={d2}))'],
             dxf=DifferentialStyle(fill=cf_fill(WINDOW_DOE), font=Font(bold=True)),
             stopIfTrue=True))
    # 3) VAEウィンドウ（DOE±2日）。AND/OR は途中で打ち切らずに全部を計算するので、
    #    DOE②が空のスロットで ""-2 が #VALUE! にならないよう N() で数値にしてから引く
    ws.conditional_formatting.add(
        rng,
        Rule(type="expression",
             formula=[f'OR(AND({d1}<>"",ABS(C$5-N({d1}))<=2),'
                      f'AND({d2}<>"",ABS(C$5-N({d2}))<=2))'],
             dxf=DifferentialStyle(fill=cf_fill(WINDOW)), stopIfTrue=True))

    # ---- 入力規則（入力時のヒントと、単位の取り違えの警告）----------------
    dv_temp = DataValidation(
        type="decimal", operator="between", formula1="34", formula2="42", allow_blank=True,
        errorStyle="warning", showErrorMessage=True, showInputMessage=True,
        errorTitle="体温の確認", error="体温は℃で入力します（例：38.5）。このまま入力しますか？",
        promptTitle="体温（℃）", prompt="黄色の列（DOE±2日）の最高・最低体温。例：38.5")
    dv_wbc = DataValidation(
        type="decimal", operator="between", formula1="1000", formula2="300000", allow_blank=True,
        errorStyle="warning", showErrorMessage=True, showInputMessage=True,
        errorTitle="WBCの単位の確認",
        error="WBCは /μL で入力します（例：12000）。×10³/μL（例：12.0）で入れると判定を誤ります。"
              "このまま入力しますか？",
        promptTitle="WBC（/μL）", prompt="/μLで入力。例：12000（≧12,000 または ≦4,000 が異常）")
    dv_abx = DataValidation(
        type="list", formula1='"はい"', allow_blank=True, showInputMessage=True,
        promptTitle="新規抗菌薬", prompt="NHSN適格抗菌薬を開始した日に「はい」")
    dv_pvap = DataValidation(
        type="list", formula1='"基準1,基準2,基準3"', allow_blank=True, showInputMessage=True,
        promptTitle="PVAP基準", prompt="該当した日に 基準1〜3 を選択")
    for dv in (dv_temp, dv_wbc, dv_abx, dv_pvap):
        ws.add_data_validation(dv)
    for n in range(1, NCLIN + 1):
        b = clin_base(n)
        dv_temp.add(f"C{b}:{gl(CN)}{b+1}")
        dv_wbc.add(f"C{b+2}:{gl(CN)}{b+3}")
        dv_abx.add(f"C{b+4}:{gl(CN)}{b+4}")
        dv_pvap.add(f"C{b+6}:{gl(CN)}{b+6}")

    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 19
    for c in range(C0, CN + 1):
        ws.column_dimensions[gl(c)].width = 6.5
    ws.freeze_panes = "C6"
    ws.sheet_view.zoomScale = 85
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_title_cols = "A:B"
    ws.print_title_rows = "4:5"
    return ws
