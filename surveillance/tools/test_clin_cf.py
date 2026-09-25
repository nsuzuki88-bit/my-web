# -*- coding: utf-8 -*-
"""臨床所見シートの色分け（条件付き書式）が正しいセルで真になるかを検証する

  ① 未割当のスロット＝灰色  ② DOE当日（①・②）＝濃い橙  ③ DOE±2日＝黄色

条件付き書式の式を、同じ座標の別シートに書き出して LibreOffice に評価させ、
スロットの DOE①・DOE② から求めた期待値と照合する（test_cf.py と同じ方式）。

  python3 test_clin_cf.py make  <テスト用ブック(make_tests.py の出力)> <出力>
  python3 test_clin_cf.py check <再計算済み(strip_charts後)>
"""
import datetime as dt
import sys, openpyxl
from openpyxl.utils import get_column_letter as gl
from common import *

DAYS = 25            # 4/1〜4/25
SLOTS = 12           # 先頭12スロット（VAC患者9名＋未割当3）
SHEETS = ("_CF_GRAY", "_CF_DOE", "_CF_WIN")


def rules(ws):
    """臨床所見の条件付き書式の式（優先順）"""
    out = []
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            out.append((str(cf.sqref), rule.formula[0]))
    return out


def make(src, dst):
    wb = openpyxl.load_workbook(src)
    cl = wb["臨床所見"]
    rs = rules(cl)
    assert len(rs) == 3, rs
    for name, (_sq, f) in zip(SHEETS, rs):
        ws = wb.create_sheet(name)
        for n in range(1, SLOTS + 1):
            b = clin_base(n)
            for off in range(8):
                for d in range(DAYS):
                    c = C0 + d
                    g = (f.replace("C$5", f"臨床所見!{gl(c)}$5")
                          .replace("$A$6:", "臨床所見!$A$6:"))
                    ws.cell(b + off, c, "=" + g)
    wb.save(dst)
    print("wrote", dst)


def check(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    cl = wb["臨床所見"]
    gray, doe, win = (wb[s] for s in SHEETS)
    ok = ng = 0

    def chk(label, got, want):
        nonlocal ok, ng
        if got == want:
            ok += 1
        else:
            ng += 1
            print(f"  NG  {label}: got {got!r} want {want!r}")

    D = lambda v: v.date() if isinstance(v, dt.datetime) else v
    for n in range(1, SLOTS + 1):
        b = clin_base(n)
        blk = cl.cell(b + 1, 1).value
        d1, d2 = D(cl.cell(b + 4, 1).value), D(cl.cell(b + 6, 1).value)
        does = [x for x in (d1, d2) if isinstance(x, dt.date)]
        rows_win, rows_doe = [], []
        for off in range(8):
            for d in range(DAYS):
                c = C0 + d
                day = D(cl.cell(5, c).value)
                want_gray = blk in (None, "")
                want_doe = day in does
                want_win = any(abs((day - x).days) <= 2 for x in does)
                lab = f"スロット{n} {gl(c)}{b + off}"
                chk(f"{lab} 灰色", gray.cell(b + off, c).value, want_gray)
                if not want_gray:
                    chk(f"{lab} DOE当日", doe.cell(b + off, c).value, want_doe)
                    chk(f"{lab} DOE±2日", win.cell(b + off, c).value, want_win)
                if off == 0:
                    if win.cell(b, c).value is True:
                        rows_win.append(day.day)
                    if doe.cell(b, c).value is True:
                        rows_doe.append(day.day)
        print(f"  スロット{n:2d} ブロック{blk!s:>4s} DOE①={d1} DOE②={d2}  "
              f"橙={rows_doe} 黄={[x for x in rows_win if x not in rows_doe]}")
    print(f"\n==== ok={ok}  NG={ng} ====")
    return ng


if __name__ == "__main__":
    if sys.argv[1] == "make":
        make(sys.argv[2], sys.argv[3])
    else:
        sys.exit(1 if check(sys.argv[2]) else 0)
