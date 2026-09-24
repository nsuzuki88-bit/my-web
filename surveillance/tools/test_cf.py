# -*- coding: utf-8 -*-
"""Allシートの条件付き書式（DOE赤／VAC基準薄赤）が正しいセルで真になるかを検証する。

条件付き書式の式そのものを、同じ行・列の別シートに書き出して LibreOffice に評価させる。
ROW()/COLUMN() は同じ座標なのでそのまま、自セル参照と日付参照だけ All! に読み替える。
"""
import sys, openpyxl
from openpyxl.utils import get_column_letter as gl
from common import *
import sheet_all

DAYS = 25                      # 4/1 〜 4/25 を検査
BLOCKS = range(1, 10)

# ブロック -> (DOEであるべきMV日のリスト, VAC基準に合致するがDOEではないMV日のリスト)
EXPECT = {
    1: ([5], []),              # PEEP経路
    2: ([], []),               # PEEP 0〜5同等 → 該当なし
    3: ([4], []),              # FiO2経路
    4: ([3], []),              # 最早DOE
    5: ([3, 17], [7]),         # 14日イベント期間：7日目は薄赤どまり
    6: ([], []),               # 第1エピソードに悪化なし
    7: ([], []),               # NIVは対象外
    8: ([4], []),              # IVAC
    9: ([4], []),              # PVAP
}


def make(src, dst):
    wb = openpyxl.load_workbook(src)
    doe_f, vac_f = sheet_all.doe_formula(), sheet_all.vac_formula()
    for name, tmpl in (("_CF_DOE", doe_f), ("_CF_VAC", vac_f)):
        ws = wb.create_sheet(name)
        for k in BLOCKS:
            for off in (1, 2):                      # FiO2行・PEEP行
                r = all_v(k) + off
                for d in range(DAYS):
                    c = C0 + d
                    f = (tmpl.replace("COUNT(C6)", f"COUNT(All!{gl(c)}{r})")
                              .replace("C$5", f"All!{gl(c)}$5"))
                    ws.cell(r, c, "=" + f)
    wb.save(dst)
    print("wrote", dst)


def verify(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    doe, vac = wb["_CF_DOE"], wb["_CF_VAC"]
    ok = ng = 0

    def chk(label, got, want):
        nonlocal ok, ng
        if got == want:
            ok += 1
        else:
            ng += 1
            print(f"  NG  {label}: got {got!r} want {want!r}")

    for k in BLOCKS:
        want_doe, want_cand = EXPECT[k]
        got_doe, got_cand = [], []
        for d in range(1, DAYS + 1):
            c, r = C0 + d - 1, all_f(k)
            dv, vv = doe.cell(r, c).value, vac.cell(r, c).value
            # FiO2行とPEEP行で同じ結果になること
            chk(f"ブロック{k} MV{d}日目 DOE式がFiO2行/PEEP行で一致",
                doe.cell(all_p(k), c).value, dv)
            chk(f"ブロック{k} MV{d}日目 VAC式がFiO2行/PEEP行で一致",
                vac.cell(all_p(k), c).value, vv)
            if dv is True:
                got_doe.append(d)
            elif vv is True:
                got_cand.append(d)
        print(f"ブロック{k}: 濃い赤(DOE)={got_doe}  薄い赤(VAC基準)={got_cand}")
        chk(f"  ブロック{k} DOE日", got_doe, want_doe)
        chk(f"  ブロック{k} VAC基準のみの日", got_cand, want_cand)

    # 判定と無関係のセルが赤くならないこと（C有無行・部屋行・空欄セル）
    for k in BLOCKS:
        for c in range(C0, C0 + DAYS):
            for r in (all_v(k), all_c(k), all_rm(k)):
                if doe.cell(r, c).value is not None:
                    chk(f"ブロック{k} {gl(c)}{r} 非対象行は評価されない", "式あり", "式なし")

    print(f"\n==== ok={ok}  NG={ng} ====")
    return ng


if __name__ == "__main__":
    if sys.argv[1] == "make":
        make(sys.argv[2], sys.argv[3])
    else:
        sys.exit(1 if verify(sys.argv[2]) else 0)
