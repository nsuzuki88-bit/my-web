# -*- coding: utf-8 -*-
"""年度サーベイランスワークシートに VAE 個別判定シートの自動作成機能を組み込む

Allシートの構造（日付行・ブロック先頭行・ブロック数）はファイルごとに違うため、
実ファイルから読み取ってから各シートを生成する。
"""
import datetime as dt
import sys, time
import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as gl
import common


def probe(ws):
    """Allシートの実構造を読み取る"""
    # ブロック先頭行＝B列が「V有無」の行
    vrows = [r for r in range(1, min(ws.max_row, 200) + 1)
             if ws.cell(r, 2).value == "V有無"]
    assert len(vrows) >= 2, "「V有無」の行が見つかりません"
    block0, bstep = vrows[0], vrows[1] - vrows[0]

    # 日付行＝ブロック先頭より上で、C列が日付になっている一番下の行
    drows = [r for r in range(1, block0)
             if isinstance(ws.cell(r, 3).value, dt.datetime)]
    assert drows, "日付の行が見つかりません"
    date_row = drows[-1]

    # 日付列の範囲
    cols = [c for c in range(1, ws.max_column + 1)
            if isinstance(ws.cell(date_row, c).value, dt.datetime)]
    c0, cn = cols[0], cols[-1]

    last = ws.max_row
    nblk = (last - block0 + 1) // bstep

    # 日ごとの集計行（集計シートの分母）＝ブロック先頭より上で、A列がこのラベルの行
    want = {"ROW_NIV": "人工呼吸器（NIV）", "ROW_V": "人工呼吸器使用患者数",
            "ROW_C": "CV挿入患者数", "ROW_ICU": "ICU入室患者数"}
    rows = {}
    for key, label in want.items():
        hit = [r for r in range(1, block0) if str(ws.cell(r, 1).value or "").strip() == label]
        assert hit, f"Allシートに「{label}」の行が見つかりません"
        rows[key] = hit[0]
    return dict(C0=c0, CN=cn, DATE_ROW=date_row, BLOCK0=block0, BSTEP=bstep,
                NBLK=nblk, ALL_LAST_ROW=last, **rows)


LABELS = ["V有無", "FiO2", "PEEP", "C有無", "部屋"]


def repair_labels(ws, g):
    """各ブロックのB列ラベルを正規の5項目に直す（例：'oo' → 'C有無'）"""
    fixed = []
    for k in range(1, g["NBLK_ALL"] + 1):
        b = g["BLOCK0"] + g["BSTEP"] * (k - 1)
        # 空のブロック（患者が入っていない）は触らない
        if all(ws.cell(b + i, 2).value in (None, "") for i in range(g["BSTEP"])):
            continue
        for i, want in enumerate(LABELS):
            got = ws.cell(b + i, 2).value
            if got != want:
                ws.cell(b + i, 2).value = want
                fixed.append((k, b + i, got, want))
    return fixed


def count_mv_blocks(ws, g):
    """MV日数が MV_MIN 日以上のブロック数を数える（判定シートの必要枚数の目安）"""
    n = 0
    for k in range(1, g["NBLK"] + 1):
        b = g["BLOCK0"] + g["BSTEP"] * (k - 1)
        mv = 0
        for c in range(g["C0"], g["CN"] + 1):
            v = ws.cell(b, c).value
            if v == "V有":
                mv += 1
            elif v in (None, ""):
                if ws.cell(b + 1, c).value is not None or ws.cell(b + 2, c).value is not None:
                    mv += 1
        if mv >= common.MV_MIN:
            n += 1
    return n


def build_guide(wb, GUIDE):
    from common import (F_TITLE, F_SEC, fill, Font, A_L, A_LW)
    ws = wb.create_sheet("使い方")
    ws.sheet_properties.tabColor = "2E7D32"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "VAE・CLABSI 自動判定ワークブックの使い方"
    ws["A1"].font = F_TITLE
    ws["A1"].fill = fill("2E7D32")
    ws.merge_cells("A1:C1")
    ws.row_dimensions[1].height = 27.75
    for i, (a, b) in enumerate(GUIDE):
        r = 3 + i
        if a and not b:
            ws.cell(r, 1, a).font = F_SEC
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        elif a:
            ws.cell(r, 2, a).font = Font(bold=True, size=9)
            ws.cell(r, 2).alignment = A_L
            ws.cell(r, 3, b).font = Font(size=9)
            ws.cell(r, 3).alignment = A_LW
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 95
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:C{2 + len(GUIDE)}"
    return ws


def guide_text():
    from common import NSHEET, NCLIN, MV_MIN, NBLK, NEV
    last = f"VAE-{NSHEET:02d}"
    return [
        ("", ""),
        ("1. 全体の流れ", ""),
        ("① Allシート",
         "これまで通り V有無／FiO2／PEEP／C有無／部屋 を日付ごとに入力します。入力方法は変わりません。"),
        ("VAE：② VAE対象一覧",
         f"Allのデータから全ブロックの VAC判定（DOE）を自動で行います（MV{MV_MIN}日以上のブロック。"
         f"イベント期間14日の重複排除まで適用）。VACの患者はセクションAにDOEの早い順に並びます。"),
        ("VAE：③ 臨床所見",
         f"VACと判定された患者だけがスロットに並びます（最大{NCLIN}名）。黄色の列（DOE±2日）に"
         f"体温・WBC・新規抗菌薬・PVAP基準を入力します。"),
        (f"VAE：④ VAE-01〜{last[-2:]}",
         f"③で臨床所見を入力した患者だけに個別判定シート（{NSHEET}枚）が割り当てられ、IVAC・PVAPまで判定されます。"),
        ("VAE：⑤ VAE集計",
         "人工呼吸器使用比と VAC／IVAC／PVAP の件数・発生率（/1,000人工呼吸器日）を月別に自動集計し、グラフにします。"),
        ("CLABSI：② CLABSI判定",
         "Allで「C有」と入力された患者がセクションBに自動で集まります。血液培養が陽性になった患者"
         f"（またはCSEPを疑う患者）だけをセクションAに1行ずつ入力すると、CLABSIかどうかを判定します（最大{NEV}件）。"),
        ("CLABSI：③ CLABSI集計",
         "中心ライン使用比と CLABSI の件数・発生率（/1,000中心ライン日。LCBI＋CSEPとLCBIのみ）を月別に自動集計し、グラフにします。"),
        ("既存の集計シート",
         "年間集計・上半期・下半期の各シートの発生件数（これまで手入力）は、⑤・③から自動で入るようにしました。"
         "既存のグラフもそのまま自動で更新されます。"),
        ("自動反映",
         "すべて数式でつながっているため、Allシートの値を直したその瞬間にすべて再計算されます。"
         "判定にマクロは使っていないので、マクロが禁止されたPCでも動きます。"),
        ("", ""),
        ("2. どこに何を入力するか", ""),
        ("Allシート", "V有無／FiO2／PEEP／C有無／部屋、A列に患者ID・氏名（従来どおり）"),
        ("VAE対象一覧 セクションB", "年齢・臓器提供同意日・備考（黄色のセル）。必要な患者だけで構いません"),
        ("臨床所見", "VACと判定された患者の 体温・WBC・新規抗菌薬・PVAP基準（黄色＝DOE±2日の列）"),
        ("VAE-01〜", "入力不要です。エピソードが2回ある患者のみ C9 を 2 に変更してください"),
        ("CLABSI判定 セクションA",
         "患者ID・血液培養採取日・検出菌名・菌の分類・症状の初日・2次性BSI・CSEP（黄色のセル）。"
         "セルを選ぶと入力のヒントが出ます"),
        ("VAE集計・CLABSI集計", "入力不要です（すべて自動計算）"),
        ("", ""),
        ("3. セルの色", ""),
        ("黄色", "入力セル（ここだけを編集します）"),
        ("薄い灰青色", "自動計算（数式。上書きしないでください）"),
        ("緑色／桃色", "判定結果（桃色はVAC・CLABSIなど該当あり）"),
        ("橙色（判定シート）", "VAEウィンドウ期間（DOE±2日）の行"),
        ("臨床所見の色分け",
         "項目ごとに 体温＝橙、WBC＝青、抗菌薬＝緑、PVAP基準＝紫。スロットの境目は太線、月の変わり目は縦の太線。"
         "黄色の列＝DOE±2日（ここに入力）、濃い橙の列＝DOE当日、灰色のスロット＝未割当（入力不要）"),
        ("", ""),
        ("3-2. Allシートの赤色（FiO2・PEEPの行）", ""),
        ("濃い赤（白文字）",
         "DOE（イベント発生日）。VAE対象一覧セクションBが判定した日で、イベント期間14日による重複排除まで適用済みです。"),
        ("薄い赤",
         "JHAISの酸素化悪化基準（VAC）に合致した日のうち、上のDOEに当たらないもの。"
         "＝直前のDOEから14日以内で新規報告の対象外か、一覧シートのスキャン範囲より後ろのブロックです。"),
        ("補足",
         "元のAllシートに入っていた条件付き書式は、#REF!エラーや参照位置のずれ"
         "（入力したセルと無関係のセルが赤くなる）があったため削除し、上の2ルールに置き換えました。"),
        ("", ""),
        ("4. MV日（人工呼吸器装着日）の判定ルール", ""),
        ("MV日とみなす日", "AllシートのV有無行が「V有」の日。V有無が空欄のときは、FiO2／PEEP に入力があればMV日とみなします。"),
        ("除外する日", "「V無」の日と「V有（NIV）」の日。非侵襲的換気（NPPV・マスクCPAP等）はVAEの分子・対象に含めません。"),
        ("エピソード", "1暦日以上完全に人工呼吸を離脱したあとの再装着は別エピソードです。"
                       "判定シートのC9「エピソード番号」を 2 に変えると2回目を判定できます。"),
        ("FiO2の単位", "小数（0.4）でも％（40）でも、どちらで入力しても同じ判定になります。"),
        ("", ""),
        ("5. VAEの判定ロジック（NHSN / JHAIS）", ""),
        ("日最小値", "その暦日内で1時間を超えて維持された最も低い PEEP／FiO2。PEEP 0〜5cmH2O は同等（5として自動補正）。"),
        ("ベースライン期間", "日最小 FiO2 または PEEP が2暦日以上安定または低下している期間。悪化初日の直前2日。"),
        ("VAC", "ベースライン初日の値と比べて FiO2 が20ポイント以上上昇、または PEEP が3cmH2O以上上昇し、"
                "それが2暦日以上持続。最も早いDOEは MV3日目。"),
        ("IVAC", "VAC成立＋MV3日目以降かつVAEウィンドウ（DOE±2日）内に、"
                 "①体温>38℃ または <36℃、あるいは WBC≥12,000 または ≤4,000/mm³ かつ "
                 "②新規抗菌薬開始（4QAD以上継続）。"),
        ("PVAP", "IVAC成立＋ウィンドウ内に PVAP基準1〜3 のいずれか。"),
        ("イベント期間", "DOEを1日目とする14日間は新たなVAEを判定しません（次の最早DOEは15日目）。"),
        ("", ""),
        ("5-2. CLABSIの判定ロジック（JHAIS Ver.2.5）", ""),
        ("中心ライン日", "Allシートの C有無 が「C有」の日。1暦日以上「C有」が途切れると、次の「C有」から数え直します。"),
        ("LCBI-1", "血液培養から認定病原体（皮膚常在菌以外）。DOE＝血液培養採取日。"),
        ("LCBI-2", "同じ皮膚常在菌が別々に採取した2セット以上から検出＋発熱(>38.0℃)・悪寒戦慄・低血圧のいずれかが"
                   "IWP（採取日±3日）内にある。DOE＝採取日と症状の初日のうち早い方。"),
        ("CSEP", "JHAISの基準。血液培養が陰性・未実施で、他に原因・他部位感染がなく、発熱・低血圧・乏尿があり、"
                 "医師が敗血症として治療を開始。DOE＝症状の初日。"),
        ("CLABSI", "上のいずれかを満たし、DOEまたはその前日に『2暦日を超えて（CL3日目以降）留置された中心ライン』がある。"
                   "抜去日と抜去翌日のDOEも含みます。"),
        ("対象外になるもの", "皮膚常在菌1セットのみ（汚染菌）、他部位感染の2次性BSI、ICU入室2日以内のDOE"
                             "（POAまたは転棟元に帰属）、ICU退室の翌々日以降、RIT内の再陽性。"),
        ("RIT（14日間）", "一次BSIのDOEから14日間は、同じ患者の血流感染を新規にしません（菌は前回に追加）。"
                          "同じ患者の血液培養は、日付の古い順に上から入力してください。"),
        ("", ""),
        ("5-3. 使用比と発生率の計算（JHAIS）", ""),
        ("発生率", "件数 ÷ 延べ医療器具使用日数 × 1,000（VAE：/1,000人工呼吸器日、CLABSI：/1,000中心ライン日）"),
        ("使用比", "延べ医療器具使用日数 ÷ 延べ入室患者日数。分母はAllシート上部の日ごとの集計行の月合計です。"),
        ("年度計", "各月の率の平均ではなく、年間の合計から計算します（pooled）。既存の年間集計シートの使用比もこの方法に直しました。"),
        ("", ""),
        ("6. 足りなくなったら", ""),
        ("判定シートを増やす", f"{last} のシートタブを右クリック →「移動またはコピー」→「コピーを作成する」。"
                               f"できたシートのC4とB78の数式内の {NSHEET} を {NSHEET+1} に書き換えます。"),
        ("スキャン範囲", f"VAE対象一覧セクションBとCLABSI判定セクションBは Allシートの先頭{NBLK}ブロックを見ています。"
                         f"それより後ろにデータがある場合は各セクションの末尾に警告が出ます。"),
        ("特定の患者に固定", "判定シートのC4に、一覧シートセクションBのブロック番号を直接入力すると固定できます。"),
        ("", ""),
        ("7. 出典・注意", ""),
        ("NHSN VAE", "CDC/NHSN Patient Safety Component Manual, Chapter 10: Ventilator-Associated Event (VAE). "
                     "https://www.cdc.gov/nhsn/pdfs/pscmanual/10-vae_final.pdf"),
        ("NHSN BSI", "CDC/NHSN Patient Safety Component Manual, Chapter 4: Bloodstream Infection Event (CLABSI). "
                     "https://www.cdc.gov/nhsn/pdfs/pscmanual/4psc_clabscurrent.pdf"),
        ("JHAIS", "日本環境感染学会 JHAIS委員会 医療器具関連感染サーベイランス部門 マニュアル Ver.2.5 "
                  "（2) 中心ライン関連血流感染、5) 人工呼吸器関連イベント、感染率と医療器具使用比の計算）"),
        ("注意", "自動判定は判定の補助です。最終判定は原典プロトコールと診療記録に基づき感染管理担当者が行ってください。"),
        ("配列数式", "VAE対象一覧の非表示列P・Q、判定シート78行目、CLABSI判定の非表示列の数式は配列数式"
                     "（数式バーで {=…} と表示）です。編集しないでください。編集してEnterだけで確定すると、判定が止まることがあります。"),
    ]


def main():
    SRC = sys.argv[1] if len(sys.argv) > 1 else "ws.xlsx"
    DST = sys.argv[2] if len(sys.argv) > 2 else "out.xlsx"
    NSHEET_ARG = int(sys.argv[3]) if len(sys.argv) > 3 else None
    NBLK_CAP = int(sys.argv[4]) if len(sys.argv) > 4 else None
    NCLIN_ARG = int(sys.argv[5]) if len(sys.argv) > 5 else None
    t0 = time.time()
    print("load", SRC)
    is_xlsm = SRC.lower().endswith(".xlsm")
    wb = openpyxl.load_workbook(SRC, keep_vba=is_xlsm)
    assert "All" in wb.sheetnames, wb.sheetnames
    al = wb["All"]

    g = probe(al)
    g["NBLK_ALL"] = g["NBLK"]
    print("Allシートの構造:", ", ".join(f"{k}={v}" for k, v in g.items()))

    fixed = repair_labels(al, g)
    if fixed:
        print(f"B列のラベルを修復: {len(fixed)}件")
        for k, row, got, want in fixed:
            print(f"  ブロック{k} {row}行目: {got!r} → {want!r}")

    g["NBLK"] = min(g["NBLK"], NBLK_CAP or common.NBLK_MAX)
    del g["NBLK_ALL"]
    common.configure(**g)

    # VAC判定される患者数から、臨床所見のスロット数と個別判定シートの枚数を決める
    import ref_vae
    _, ref = ref_vae.analyse(SRC)
    mv_ok = [r for r in ref if r["mvtotal"] >= common.MV_MIN]
    vac = [r for r in mv_ok if r["doe"]]
    nclin = NCLIN_ARG or max(40, ((len(vac) * 3 + 9) // 10) * 10)
    nsheet = NSHEET_ARG or common.NSHEET          # 既定10枚（足りなければ第3引数で増やす）
    common.configure(NCLIN=nclin, NSHEET=nsheet)
    print(f"MV{common.MV_MIN}日以上: {len(mv_ok)}名 / うちVAC判定: {len(vac)}名"
          f" → 臨床所見スロット{nclin}件・個別判定シート{nsheet}枚")
    print(f"一覧セクションBのスキャン範囲: 先頭{g['NBLK']}ブロック")

    # configure() の後に import すること（各モジュールは import 時に定数を取り込む）
    import sheet_clin, sheet_vae, sheet_list, sheet_all, sheet_clabsi, sheet_summary
    from common import (N_ALL, N_ALLD, N_CLIN, CN, ALL_LAST_ROW,
                        CLIN_LAST_ROW, NSHEET)

    from common import N_D0, N_D1, N_D2, N_DP, C0
    for nm, ref in ((N_ALL,  f"All!$A$1:${gl(CN)}${ALL_LAST_ROW}"),
                    (N_ALLD, f"All!$C$1:${gl(CN)}${ALL_LAST_ROW}"),
                    (N_CLIN, f"臨床所見!$A$1:${gl(CN)}${CLIN_LAST_ROW}"),
                    # VAC判定用：当日・前日・前々日・翌日の4本。
                    # A列・B列のラベル文字を拾わないよう、当日はE列から始める
                    (N_D0, f"All!${gl(C0+2)}$1:${gl(CN)}${ALL_LAST_ROW}"),
                    (N_D1, f"All!${gl(C0+1)}$1:${gl(CN-1)}${ALL_LAST_ROW}"),
                    (N_D2, f"All!${gl(C0)}$1:${gl(CN-2)}${ALL_LAST_ROW}"),
                    (N_DP, f"All!${gl(C0+3)}$1:${gl(CN+1)}${ALL_LAST_ROW}")):
        if nm in wb.defined_names:
            del wb.defined_names[nm]
        wb.defined_names.add(DefinedName(nm, attr_text=ref))

    removed = sheet_all.build(wb)
    print(f"All: 壊れていた条件付き書式 {removed} ルールを削除し、DOE強調の2ルールに置換")
    build_guide(wb, guide_text())
    sheet_clin.build(wb)
    sheet_list.build(wb)
    for n in range(1, NSHEET + 1):
        sheet_vae.build(wb, n)
    sheet_summary.build_vae(wb)
    sheet_clabsi.build(wb)
    sheet_summary.build_clabsi(wb)
    for msg in sheet_summary.wire_legacy(wb):
        print("既存シート:", msg)
    print("sheets built", round(time.time() - t0, 1), "s")

    # Excelで範囲全体を計算させるため、配列評価の数式は配列数式で保存する
    new_sheets = (["VAE対象一覧", "臨床所見"] + [f"VAE-{n:02d}" for n in range(1, NSHEET + 1)]
                  + ["VAE集計", "CLABSI判定", "CLABSI集計"])
    narr = sum(common.as_array_formulas(wb[s]) for s in new_sheets)
    print(f"配列数式に変換: {narr}セル")

    order = (["使い方", "All", "VAE対象一覧", "臨床所見"]
             + [f"VAE-{n:02d}" for n in range(1, NSHEET + 1)]
             + ["VAE集計", "CLABSI判定", "CLABSI集計"])
    rest = [s for s in wb.sheetnames if s not in order]
    wb._sheets = [wb[s] for s in order + rest]

    wb.calculation.fullCalcOnLoad = True
    wb.active = 0

    print("save", DST)
    wb.save(DST)
    print("done", round(time.time() - t0, 1), "s")


if __name__ == "__main__":
    main()
