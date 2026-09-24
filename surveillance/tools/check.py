# -*- coding: utf-8 -*-
"""再計算済みブックを読み、NHSN/JHAISの期待値と照合する"""
import datetime as dt
import openpyxl
from common import *

D = lambda m, d: dt.datetime(2025, m, d)

# ブロック -> (MV日数, エピソード総数, DOE1, 判定1, DOE2, 判定2)
EXPECT_BLOCK = {
    1: (6,  1, D(4, 5),  "VAC",  None, ""),
    2: (6,  1, None,     "",     None, ""),
    3: (6,  1, D(4, 4),  "VAC",  None, ""),
    4: (4,  1, D(4, 3),  "VAC",  None, ""),
    5: (20, 1, D(4, 3),  "VAC",  D(4, 17), "VAC"),
    6: (13, 2, None,     "",     None, ""),
    7: (0,  0, None,     "",     None, ""),
    8: (6,  1, D(4, 4),  "IVAC", None, ""),
    9: (6,  1, D(4, 4),  "PVAP", None, ""),
    10: (4, 1, D(4, 3),  "VAC",  None, ""),
    11: (4, 1, D(4, 3),  "VAC",  None, ""),
    12: (4, 1, D(4, 3),  "VAC",  None, ""),
}
# ブロック -> 期待する割当シート番号（Noneは未割当）
# MV日数が4日以上のブロックだけに、上から順に判定シートが割り当てられる
EXPECT_SLOT = {1: 1, 2: 2, 3: 3, 4: 4, 5: 5, 6: 6, 7: None, 8: 7, 9: 8,
               10: 9, 11: 10, 12: 11}

ok, ng = 0, 0


def chk(label, got, want):
    global ok, ng
    if isinstance(want, dt.datetime) and isinstance(got, dt.datetime):
        good = got.date() == want.date()
    else:
        good = (got or "") == (want if want is not None else "") or got == want
    if good:
        ok += 1
        print(f"  ok   {label}: {got!r}")
    else:
        ng += 1
        print(f"  NG   {label}: got {got!r}  want {want!r}")


def main(path="test_calc.xlsx"):
    wb = openpyxl.load_workbook(path, data_only=True)
    lst, al = wb["VAE対象一覧"], wb["All"]

    print("=== セクションB：全ブロックスキャン ===")
    for k, slot in EXPECT_SLOT.items():
        r = LIST_B_TOP + k - 1
        mvdays = lst.cell(r, 4).value
        assigned = lst.cell(r, 7).value
        print(f"ブロック{k}  {lst.cell(r,10).value}")
        chk(f"  MV日数", mvdays, EXPECT_BLOCK[k][0])
        chk(f"  割当No", assigned, slot)

    print("\n=== セクションA／個別判定シート ===")
    for k, slot in EXPECT_SLOT.items():
        if slot is None:
            continue
        want = EXPECT_BLOCK[k]
        ws = wb[f"VAE-{slot:02d}"]
        print(f"VAE-{slot:02d} (= ブロック{k})")
        chk("  ブロック番号", ws["C4"].value, k)
        chk("  患者ID", ws["C5"].value, f"T{k:02d}")
        chk("  エピソード総数", ws["K6"].value, want[1])
        chk("  DOE①", ws["K8"].value, want[2])
        chk("  判定①", ws["K9"].value, want[3])
        chk("  DOE②", ws["K10"].value, want[4])
        chk("  判定②", ws["K11"].value, want[5])

    print("\n=== エピソード分割（ブロック6 / VAE-06）===")
    ws6 = wb["VAE-06"]
    chk("  第1エピソード MV1日目", ws6["C10"].value, D(4, 1))
    chk("  第1エピソード 日数", ws6["K5"].value, 5)
    chk("  MV日数（全期間）", ws6["C11"].value, 13)

    print("\n=== 分母（Allシート集計行）===")
    # 4/1にV有なのは 1,2,3,4,5,8,9,10,11,12 の10ブロック
    # （ブロック6はV有無が空欄、ブロック7はNIV）
    chk("  4/1 人工呼吸器使用患者数", al.cell(3 if al["B1"].value == "月日" else 2, 3).value, 10)

    print("\n=== 月別自動集計（セクションC）===")
    r_vac, r_ivac, r_pvap = LIST_C_TOP, LIST_C_TOP + 1, LIST_C_TOP + 2
    chk("  4月 VAC件数", lst.cell(r_vac, 2).value, 8)
    chk("  4月 IVAC件数", lst.cell(r_ivac, 2).value, 1)
    chk("  4月 PVAP件数", lst.cell(r_pvap, 2).value, 1)

    print(f"\n==== ok={ok}  NG={ng} ====")
    return ng


if __name__ == "__main__":
    import sys
    sys.exit(1 if main(*sys.argv[1:]) else 0)
