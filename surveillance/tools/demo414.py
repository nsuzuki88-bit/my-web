# -*- coding: utf-8 -*-
"""実データ（ブロック414・DOE 2026/9/18）で臨床所見の入力挙動を確認する"""
import datetime as dt, sys, openpyxl
from openpyxl.utils import get_column_letter as gl
from common import LIST_A_TOP, clin_base, C0, CN, DATE_ROW

T_HI, T_LO, W_HI, W_LO, ABX, ABX_MEMO, PVAP, SPEC = range(8)


def col_of(al, d):
    for c in range(C0, CN + 1):
        v = al.cell(DATE_ROW, c).value
        if isinstance(v, dt.datetime) and v.date() == d.date():
            return c
    raise ValueError(d)


def step1(src, dst):
    """体温とWBCだけを入力する（新規抗菌薬は入れない）"""
    wb = openpyxl.load_workbook(src, keep_vba=True)
    al, cl = wb["All"], wb["臨床所見"]
    b, c = clin_base(1), col_of(al, dt.datetime(2026, 9, 18))
    cl.cell(b + T_HI, c, 38.6)          # DOE当日 9/18 最高体温
    cl.cell(b + T_LO, c, 37.1)          #          最低体温
    cl.cell(b + W_HI, c + 1, 15200)     # 9/19 WBC最高
    cl.cell(b + W_LO, c + 1, 9800)      #      WBC最低
    wb.save(dst); print("step1:", dst, "（体温・WBCのみ）")


def step2(src, dst):
    """新規抗菌薬を追加する"""
    wb = openpyxl.load_workbook(src, keep_vba=True)
    al, cl = wb["All"], wb["臨床所見"]
    b, c = clin_base(1), col_of(al, dt.datetime(2026, 9, 18))
    cl.cell(b + ABX, c + 1, "はい")
    cl.cell(b + ABX_MEMO, c + 1, "MEPM 9/19-9/22（4QAD）")
    wb.save(dst); print("step2:", dst, "（＋新規抗菌薬）")


def show(path, title):
    wb = openpyxl.load_workbook(path, data_only=True)
    lst, w = wb["VAE対象一覧"], wb["VAE-01"]
    r = LIST_A_TOP
    print(f"\n=== {title} ===")
    print(f"  一覧セクションA: ブロック{lst.cell(r,2).value} "
          f"{lst.cell(r,3).value} {lst.cell(r,4).value} / "
          f"DOE {lst.cell(r,8).value:%Y/%m/%d} / 判定 {lst.cell(r,9).value} / "
          f"臨床所見 {lst.cell(r,12).value} / 判定シート {lst.cell(r,13).value}")
    print(f"\n  VAE-01（判定サマリー）: 検出VAE件数={w['K7'].value} "
          f"DOE={w['K8'].value:%Y/%m/%d} 判定={w['K9'].value}")
    hdr = ("MV日", "日付", "PEEP", "FiO2", "最高体温", "最低体温", "WBC最高",
           "WBC最低", "新規抗菌薬", "体温/WBC異常", "VAC基準", "IVAC", "判定")
    cols = (1, 2, 3, 4, 5, 6, 7, 8, 9, 14, 19, 21, 23)
    print("\n  " + "".join(f"{h:<12}" for h in hdr))
    print("  " + "-" * 145)
    for d in range(1, 11):
        row = 14 + d
        vals = []
        for c in cols:
            v = w.cell(row, c).value
            if isinstance(v, dt.datetime):
                v = f"{v:%m/%d}"
            vals.append("―" if v in (None, "") else str(v))
        mark = " ←DOE" if w.cell(row, 20).value == 1 else ""
        print("  " + "".join(f"{v:<12}" for v in vals) + mark)


if __name__ == "__main__":
    if sys.argv[1] == "step1":
        step1(sys.argv[2], sys.argv[3])
    elif sys.argv[1] == "step2":
        step2(sys.argv[2], sys.argv[3])
    else:
        show(sys.argv[2], sys.argv[3])
