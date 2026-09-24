# -*- coding: utf-8 -*-
"""NHSN/JHAIS の教科書例をAllシートに注入してテスト用ブックを作る"""
import openpyxl
from openpyxl.utils import get_column_letter as gl
from common import *

# (ブロック, 説明, V有無の書き方, FiO2リスト, PEEPリスト)  ※Noneはその日空欄
CASES = [
    (1, "PEEP経路のVAC（JHAIS 例1）", "V有",
     [40, 40, 40, 40, 40, 40], [5, 5, 5, 5, 8, 8]),
    (2, "PEEP 0〜5は同等→VACにならない", "V有",
     [40, 40, 40, 40, 40, 40], [0, 0, 0, 3, 3, 3]),
    (3, "FiO2経路のVAC（+20ポイント）", "V有",
     [40, 40, 40, 60, 60, 60], [5, 5, 5, 5, 5, 5]),
    (4, "最早DOEはMV3日目", "V有",
     [40, 40, 70, 70], [5, 5, 5, 5]),
    (5, "イベント期間14日（2件目は15日目以降）", "V有",
     [40, 40, 70, 70, 40, 40, 70, 70, 40, 40,
      40, 40, 40, 40, 40, 40, 70, 70, 40, 40], [5] * 20),
    (6, "エピソード分割（1暦日以上の離脱）", None,
     [40, 40, 40, 40, 40, None, None, 40, 40, 40, 40, 40, 40, 40, 40],
     [5, 5, 5, 5, 5, None, None, 5, 5, 5, 5, 5, 5, 5, 5]),
    (7, "NIVは対象外", "V有（NIV）",
     [40, 40, 40, 60, 60, 60], [5, 5, 5, 5, 5, 5]),
    (8, "IVAC（体温異常＋新規抗菌薬）", "V有",
     [40, 40, 40, 60, 60, 60], [5, 5, 5, 5, 5, 5]),
    (9, "PVAP（IVAC＋基準1）", "V有",
     [40, 40, 40, 60, 60, 60], [5, 5, 5, 5, 5, 5]),
]

# 臨床所見： (ブロック, MV日, 行オフセット, 値)
CLIN = [
    (8, 4, 0, 38.5),        # 最高体温 38.5℃
    (8, 5, 4, "はい"),      # 新規抗菌薬開始
    (9, 4, 0, 38.5),
    (9, 5, 4, "はい"),
    (9, 5, 6, "基準1"),     # PVAP基準1
]


def main(src="out.xlsx", dst="test.xlsx"):
    wb = openpyxl.load_workbook(src)
    al, cl, lst = wb["All"], wb["臨床所見"], wb["VAE対象一覧"]

    for k, desc, vmark, fio2, peep in CASES:
        al.cell(all_v(k), 1, f"T{k:02d}")           # 患者ID
        al.cell(all_p(k), 1, f"検証{k:02d}")        # 氏名
        for d, (f, p) in enumerate(zip(fio2, peep)):
            c = C0 + d
            if f is None and p is None:
                continue
            if vmark:
                al.cell(all_v(k), c, vmark)
            al.cell(all_f(k), c, f)
            al.cell(all_p(k), c, p)
            al.cell(all_rm(k), c, "A-1")
        lst.cell(LIST_B_TOP + k - 1, 8, 70)          # 年齢
        lst.cell(LIST_B_TOP + k - 1, 10, desc)       # 備考

    for k, day, off, val in CLIN:
        cl.cell(clin_base(k) + off, C0 + day - 1, val)

    wb.save(dst)
    print("wrote", dst)


if __name__ == "__main__":
    main()
