# -*- coding: utf-8 -*-
"""VAE集計・CLABSI集計シート（使用比と発生率の月別自動集計とグラフ）

  分母 … Allシート上部の日ごとの集計行（「ICU入室患者数」「人工呼吸器使用患者数」
         「CV挿入患者数」）を、日付行で月ごとに合計する（SUMIFS）。
  分子 … VAE：VAE対象一覧セクションAの DOE と判定（VAC／IVAC／PVAP）
         CLABSI：CLABSI判定シートの DOE・基準・判定
  計算式はJHAISマニュアルVer.2.5：
         感染率＝件数÷延べ医療器具使用日数×1000、医療器具使用比＝延べ医療器具使用日数÷延べ入室患者日数
         年度計は月ごとの率の平均ではなく、年間の合計から求める（pooled）。

既存の上半期／下半期／年間集計シートの発生件数（手入力だった行）も、ここから自動で入るようにつなぐ。
"""
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (CharacterProperties, Paragraph, ParagraphProperties,
                                   RegularTextRun)
from openpyxl.chart.title import Title
from openpyxl.chart.text import Text
from openpyxl.formatting.rule import Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter as gl
from common import *

VAE_SHEET = "VAE集計"
CL_SHEET = "CLABSI集計"

# グラフの色（dataviz の検証スクリプトで確認済み）
C_ONE = "2A78D6"                           # 1系列（使用比の折れ線）
TIERS = ["86B6EF", "2A78D6", "104281"]     # VAC→IVAC→PVAP：段階なので同じ青の濃淡（ordinal）
CATS = ["2A78D6", "EB6834"]                # LCBI・CSEP：種類なので別の色（categorical）
INK, INK2, GRID, SURFACE = "262626", "595959", "E3E3E3", "FFFFFF"

HDR = 4            # 月の見出し行
R0 = 5             # 1つ目の指標の行
CH1, CH2 = 21, 38  # グラフの位置（行）
DATA = 56          # グラフ用データの見出し行


# ---------------------------------------------------------------- 数式
def _start():
    return f"All!$C${DATE_ROW}"


def _dates():
    return f"All!$C${DATE_ROW}:${gl(CN)}${DATE_ROW}"


def _mon(i):
    """月 i（0＝年度の最初の月）の範囲条件"""
    return f'">="&EDATE({_start()},{i}),', f'"<"&EDATE({_start()},{i + 1})'


def dev_days(row, i):
    """Allシートの日ごとの集計行 row を、月 i について合計する"""
    lo, hi = _mon(i)
    d = _dates()
    return f"=SUMIFS(All!$C${row}:${gl(CN)}${row},{d},{lo}{d},{hi})"


def _count_in(rng_date, i):
    lo, hi = _mon(i)
    return f"{rng_date},{lo}{rng_date},{hi}"


# ---------------------------------------------------------------- 表
def _table(ws, title, color, rows, notes):
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws["A1"].fill = fill(color)
    ws.merge_cells("A1:N1")
    ws.row_dimensions[1].height = 27.75
    ws["A2"] = ("Allシートの入力と、各判定シートの判定結果から自動で集計します（入力するセルはありません）。"
                "月は年度の最初の月から順に並びます。まだデータのない月は空欄です。")
    ws["A2"].font = F_NOTE
    ws.merge_cells("A2:N2")

    style_cell(ws.cell(HDR, 1, "指標"), NAVY, F_HDR, A_C)
    for i in range(12):
        style_cell(ws.cell(HDR, i + 2, f'=TEXT(EDATE({_start()},{i}),"m月")'), NAVY, F_HDR, A_C)
    style_cell(ws.cell(HDR, 14, "年度計"), NAVY, F_HDR, A_C)

    for j, (label, month_f, year_f, fmt, band) in enumerate(rows):
        r = R0 + j
        style_cell(ws.cell(r, 1, label), band, Font(size=9, bold=(band == GREEN)), A_L)
        for i in range(12):
            style_cell(ws.cell(r, i + 2, month_f(i, gl(i + 2), r)), band if band == GREEN else AUTO,
                       F_RESULT if band == GREEN else F_BODY, A_C, numfmt=fmt)
        style_cell(ws.cell(r, 14, year_f(r)), band if band == GREEN else AUTO,
                   F_RESULT if band == GREEN else Font(bold=True, size=9), A_C, numfmt=fmt)
    for i, text in enumerate(notes):
        r = R0 + len(rows) + 1 + i
        c = ws.cell(r, 1, text)
        c.font = F_NOTE
        c.alignment = A_LW
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=14)
        ws.row_dimensions[r].height = 27
    ws.column_dimensions["A"].width = 34
    for i in range(2, 15):
        ws.column_dimensions[gl(i)].width = 8.5
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _chart_data(ws, series):
    """グラフ用のデータ行。データのない月は #N/A にして、グラフで0に落ちないようにする"""
    ws.cell(DATA, 1, "グラフ用データ（自動計算。編集しないでください。データのない月は #N/A＝グラフに描かない）")
    ws.cell(DATA, 1).font = Font(size=8, color="7F7F7F")
    ws.merge_cells(start_row=DATA, start_column=1, end_row=DATA, end_column=14)
    ws.cell(DATA + 1, 1, "月").font = Font(size=8, color="7F7F7F")
    for i in range(12):
        c = ws.cell(DATA + 1, i + 2, f"={gl(i + 2)}${HDR}")
        c.font = Font(size=8, color="7F7F7F")
    for j, (name, f) in enumerate(series):
        r = DATA + 2 + j
        ws.cell(r, 1, name).font = Font(size=8, color="7F7F7F")
        for i in range(12):
            c = ws.cell(r, i + 2, f(gl(i + 2)))
            c.font = Font(size=8, color="7F7F7F")
            c.number_format = "0.00"
    ws.conditional_formatting.add(
        f"B{DATA + 2}:M{DATA + 1 + len(series)}",
        Rule(type="expression", formula=[f"ISNA(B{DATA + 2})"],
             dxf=DifferentialStyle(font=Font(color="D9D9D9")), stopIfTrue=False))


# ---------------------------------------------------------------- グラフ
def _rich(size, bold=False, color=INK2):
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(
        sz=size * 100, b=bold, solidFill=color)), endParaRPr=CharacterProperties())])


def _title(text, size=11):
    cp = CharacterProperties(sz=size * 100, b=True, solidFill=INK)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def _axes(ch, ytitle, yfmt):
    hair = GraphicalProperties(ln=LineProperties(solidFill=GRID, w=6350))
    ch.y_axis.majorGridlines = ChartLines(spPr=hair)
    ch.y_axis.scaling.min = 0
    ch.y_axis.number_format = yfmt
    ch.y_axis.title = ytitle
    ch.y_axis.title.tx.rich.p[0].pPr = ParagraphProperties(
        defRPr=CharacterProperties(sz=900, b=False, solidFill=INK2))
    for ax in (ch.x_axis, ch.y_axis):
        ax.delete = False                       # openpyxl 3.1 の既定では軸が消えるため明示する
        ax.txPr = _rich(9)
        ax.spPr = GraphicalProperties(ln=LineProperties(solidFill="BFBFBF", w=6350))
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.x_axis.majorTickMark = "none"
    ch.y_axis.majorTickMark = "none"


def _size(ch):
    ch.width = 24
    ch.height = 8


def _text_categories(ws, ch):
    """横軸（月）を文字の参照にする。openpyxl の set_categories は数値の参照（numRef）で書くため、
    キャッシュで表示するプレビューや保護ビューで「4月」ではなく 1,2,3… と出てしまう"""
    from openpyxl.chart.data_source import AxDataSource, StrRef
    ref = f"'{ws.title}'!$B${DATA + 1}:$M${DATA + 1}"
    for s in ch.series:
        s.cat = AxDataSource(strRef=StrRef(f=ref))


def line_chart(ws, row, title, ytitle):
    ch = LineChart()
    ch.title = _title(title)
    ch.style = 2
    data = Reference(ws, min_col=1, max_col=13, min_row=row, max_row=row)
    ch.add_data(data, from_rows=True, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, max_col=13, min_row=DATA + 1))
    s = ch.series[0]
    s.smooth = False
    s.graphicalProperties.line.solidFill = C_ONE
    s.graphicalProperties.line.width = 22225           # 1.75pt
    s.marker.symbol = "circle"
    s.marker.size = 7
    s.marker.graphicalProperties = GraphicalProperties(solidFill=C_ONE)
    s.marker.graphicalProperties.line.solidFill = SURFACE  # マーカーの白い縁（重なりの区切り）
    _axes(ch, ytitle, "0.00")
    ch.legend = None                                     # 1系列なので凡例は不要（題名が名前）
    ch.display_blanks = "gap"
    _text_categories(ws, ch)
    _size(ch)
    return ch


def stacked_chart(ws, rows, colors, title, ytitle):
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.gapWidth = 70
    ch.title = _title(title)
    ch.style = 2
    for r in rows:
        ch.add_data(Reference(ws, min_col=1, max_col=13, min_row=r, max_row=r),
                    from_rows=True, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, max_col=13, min_row=DATA + 1))
    for s, color in zip(ch.series, colors):
        s.graphicalProperties.solidFill = color
        s.graphicalProperties.line.solidFill = SURFACE   # 積み上げの間に白い隙間
        s.graphicalProperties.line.width = 12700
    _axes(ch, ytitle, "0.0")
    ch.legend.position = "b"
    ch.legend.txPr = _rich(9)
    _text_categories(ws, ch)
    _size(ch)
    return ch


# ---------------------------------------------------------------- VAE集計
def build_vae(wb):
    ws = wb.create_sheet(VAE_SHEET)
    ws.sheet_properties.tabColor = "1F4E79"
    L = "'VAE対象一覧'"
    D1, J1 = f"{L}!$H${LIST_A_TOP}:$H${LIST_A_BOT}", f"{L}!$I${LIST_A_TOP}:$I${LIST_A_BOT}"
    D2, J2 = f"{L}!$J${LIST_A_TOP}:$J${LIST_A_BOT}", f"{L}!$K${LIST_A_TOP}:$K${LIST_A_BOT}"

    def count(kind):
        return lambda i, c, r: (f'=COUNTIFS({_count_in(D1, i)},{J1},"{kind}")'
                                f'+COUNTIFS({_count_in(D2, i)},{J2},"{kind}")')

    s = lambda r: f"=SUM(B{r}:M{r})"
    rate = lambda num: (lambda i, c, r: f'=IF({c}{R0+1}=0,"",{num(c)}/{c}{R0+1}*1000)')
    rate_y = lambda num: (lambda r: f'=IF(N{R0+1}=0,"",{num("N")}/N{R0+1}*1000)')
    R = R0
    rows = [
        ("延べ入室患者日数", lambda i, c, r: dev_days(ROW_ICU, i), s, "#,##0", None),
        ("延べ人工呼吸器使用日数", lambda i, c, r: dev_days(ROW_V, i), s, "#,##0", None),
        ("人工呼吸器使用比（使用日数÷患者日数）", lambda i, c, r: f'=IF({c}{R}=0,"",{c}{R+1}/{c}{R})',
         lambda r: f'=IF(N{R}=0,"",N{R+1}/N{R})', "0.00", GREEN),
        ("VAC 件数", count("VAC"), s, "0", None),
        ("IVAC 件数", count("IVAC"), s, "0", None),
        ("PVAP 件数", count("PVAP"), s, "0", None),
        ("VAE 合計（VAC＋IVAC＋PVAP）", lambda i, c, r: f"=SUM({c}{R+3}:{c}{R+5})", s, "0", None),
        ("VAC 発生率（/1,000人工呼吸器日）", rate(lambda c: f"{c}{R+3}"),
         rate_y(lambda c: f"{c}{R+3}"), "0.0", None),
        ("IVAC 発生率（/1,000人工呼吸器日）", rate(lambda c: f"{c}{R+4}"),
         rate_y(lambda c: f"{c}{R+4}"), "0.0", None),
        ("PVAP 発生率（/1,000人工呼吸器日）", rate(lambda c: f"{c}{R+5}"),
         rate_y(lambda c: f"{c}{R+5}"), "0.0", None),
        ("VAE 発生率 合計（/1,000人工呼吸器日）", rate(lambda c: f"{c}{R+6}"),
         rate_y(lambda c: f"{c}{R+6}"), "0.0", GREEN),
        ("IVAC-plus 発生率（IVAC＋PVAP）", rate(lambda c: f"({c}{R+4}+{c}{R+5})"),
         rate_y(lambda c: f"({c}{R+4}+{c}{R+5})"), "0.0", None),
    ]
    notes = [
        "計算式（JHAIS マニュアル Ver.2.5）：発生率＝件数÷延べ人工呼吸器使用日数×1,000、"
        "使用比＝延べ人工呼吸器使用日数÷延べ入室患者日数。年度計は各月の率の平均ではなく、年間の合計から求めます。"
        "件数は1件を最上位の段階（VAC→IVAC→PVAP）で1回だけ数えます。分母はAllシート上部の"
        "「人工呼吸器使用患者数」（V有）と「ICU入室患者数」（C有＋C無）の日ごとの値の月合計です。",
        "参考値：日本の18 ICU（2020–2022年）で VAE 10.0／1,000人工呼吸器日（Nakahashi S, et al. "
        "Intensive Care Med 2025）、中国の多施設ICUで VAC 16.7・IVAC 6.4・PVAP 1.64（He Q, et al. "
        "Crit Care 2021）。発生率は施設・患者層・実装で大きく変わるため、比較は同じ定義で行ってください。",
    ]
    _table(ws, "VAE 使用比・発生率（月別・自動集計）", "1F4E79", rows, notes)
    _chart_data(ws, [
        ("人工呼吸器使用比", lambda c: f"=IF({c}{R}=0,NA(),{c}{R+2})"),
        ("VAC", lambda c: f"=IF({c}{R+1}=0,NA(),{c}{R+7})"),
        ("IVAC", lambda c: f"=IF({c}{R+1}=0,NA(),{c}{R+8})"),
        ("PVAP", lambda c: f"=IF({c}{R+1}=0,NA(),{c}{R+9})"),
    ])
    ws.add_chart(line_chart(ws, DATA + 2, "人工呼吸器使用比（延べ人工呼吸器使用日数÷延べ入室患者日数）",
                            "使用比"), f"A{CH1}")
    ws.add_chart(stacked_chart(ws, [DATA + 3, DATA + 4, DATA + 5], TIERS,
                               "VAE 発生率（/1,000人工呼吸器日）　積み上げの高さ＝VAE合計",
                               "件／1,000人工呼吸器日"), f"A{CH2}")
    return ws


# ---------------------------------------------------------------- CLABSI集計
def build_clabsi(wb):
    ws = wb.create_sheet(CL_SHEET)
    ws.sheet_properties.tabColor = "7030A0"
    S = "'CLABSI判定'"
    col = lambda c: f"{S}!${c}${CL_EV_TOP}:${c}${CL_EV_BOT}"
    DOE, CRIT, JUDGE = col("L"), col("K"), col("O")

    def count(*crit):
        return lambda i, c, r: "=" + "+".join(
            f'COUNTIFS({JUDGE},"CLABSI",{CRIT},"{k}",{_count_in(DOE, i)})' for k in crit)

    s = lambda r: f"=SUM(B{r}:M{r})"
    R = R0
    rate = lambda num: (lambda i, c, r: f'=IF({c}{R+1}=0,"",{num(c)}/{c}{R+1}*1000)')
    rate_y = lambda num: (lambda r: f'=IF(N{R+1}=0,"",{num("N")}/N{R+1}*1000)')
    rows = [
        ("延べ入室患者日数", lambda i, c, r: dev_days(ROW_ICU, i), s, "#,##0", None),
        ("延べ中心ライン使用日数", lambda i, c, r: dev_days(ROW_C, i), s, "#,##0", None),
        ("中心ライン使用比（使用日数÷患者日数）", lambda i, c, r: f'=IF({c}{R}=0,"",{c}{R+1}/{c}{R})',
         lambda r: f'=IF(N{R}=0,"",N{R+1}/N{R})', "0.00", GREEN),
        ("CLABSI 件数（LCBI）", count("LCBI-1", "LCBI-2"), s, "0", None),
        ("CLABSI 件数（CSEP）", count("CSEP"), s, "0", None),
        ("CLABSI 件数 合計（LCBI＋CSEP）", lambda i, c, r: f"={c}{R+3}+{c}{R+4}", s, "0", None),
        ("CLABSI 発生率 合計（/1,000中心ライン日）", rate(lambda c: f"{c}{R+5}"),
         rate_y(lambda c: f"{c}{R+5}"), "0.0", GREEN),
        ("CLABSI 発生率 LCBIのみ（/1,000中心ライン日）", rate(lambda c: f"{c}{R+3}"),
         rate_y(lambda c: f"{c}{R+3}"), "0.0", None),
    ]
    notes = [
        "計算式（JHAIS マニュアル Ver.2.5）：CLABSI発生率＝CLABSI件数÷延べ中心ライン使用日数×1,000、"
        "中心ライン使用比＝延べ中心ライン使用日数÷延べ入室患者日数。年度計は年間の合計から求めます。"
        "JHAISはCLABSIを「LCBI＋CSEP」と「LCBIのみ」の両方で集計するため、両方を示します。"
        "分母はAllシート上部の「CV挿入患者数」（C有）と「ICU入室患者数」（C有＋C無）の日ごとの値の月合計です。",
        "件数は「CLABSI判定」シートで判定が「CLABSI」になった行を、DOEの月で数えます。"
        "サーベイランスの結果を現場へ定期的に返すこと（フィードバック）はCLABSIの減少と関連します"
        "（Marsteller JA, et al. Crit Care Med 2012、Ben-David D, et al. Euro Surveill 2023）。",
    ]
    _table(ws, "CLABSI 使用比・発生率（月別・自動集計）", "7030A0", rows, notes)
    _chart_data(ws, [
        ("中心ライン使用比", lambda c: f"=IF({c}{R}=0,NA(),{c}{R+2})"),
        ("LCBI", lambda c: f"=IF({c}{R+1}=0,NA(),{c}{R+3}/{c}{R+1}*1000)"),
        ("CSEP", lambda c: f"=IF({c}{R+1}=0,NA(),{c}{R+4}/{c}{R+1}*1000)"),
    ])
    ws.add_chart(line_chart(ws, DATA + 2, "中心ライン使用比（延べ中心ライン使用日数÷延べ入室患者日数）",
                            "使用比"), f"A{CH1}")
    ws.add_chart(stacked_chart(ws, [DATA + 3, DATA + 4], CATS,
                               "CLABSI 発生率（/1,000中心ライン日）　積み上げの高さ＝LCBI＋CSEP",
                               "件／1,000中心ライン日"), f"A{CH2}")
    return ws


# ---------------------------------------------------------------- 既存シートとの接続
def wire_legacy(wb):
    """既存の年間集計シートの手入力だった件数を、新しい集計シートから自動で入れる"""
    done = []
    if "年間集計(VAE)" in wb.sheetnames:
        ws = wb["年間集計(VAE)"]
        labels = [str(ws.cell(r, 1).value or "") for r in (7, 8, 9)]
        if "VAC" in labels[0] and "IVAC" in labels[1] and "PVAP" in labels[2]:
            for i in range(12):
                c = gl(i + 2)
                for r_old, r_new in ((7, R0 + 3), (8, R0 + 4), (9, R0 + 5)):
                    ws[f"{c}{r_old}"] = f"='{VAE_SHEET}'!{c}{r_new}"
            ws["N6"] = '=IF(N5=0,"",N4/N5)'
            done.append("年間集計(VAE)：VAC・IVAC・PVAP件数を自動化、年度の使用比を合計から計算")
    if "年間集計(CLABSI)" in wb.sheetnames:
        ws = wb["年間集計(CLABSI)"]
        if "発生件数" in str(ws["A7"].value or ""):
            for i in range(12):
                c = gl(i + 2)
                ws[f"{c}7"] = f"='{CL_SHEET}'!{c}{R0 + 5}"
            ws["N6"] = '=IF(N5=0,"",N4/N5)'
            if str(ws["A1"].value or "").startswith("VAE"):
                ws["A1"] = "CLABSIサーベイランス"
            done.append("年間集計(CLABSI)：件数を自動化、題名の誤り（VAE）を修正、年度の使用比を合計から計算")
            # 参照が #REF! になっていた発生密度率のグラフを直す
            for ch in ws._charts:
                refs = [s.val.numRef.f for s in ch.series if s.val and s.val.numRef]
                if refs and all("#REF!" in f for f in refs):
                    s0 = ch.series[0]
                    s0.val.numRef.f = "'年間集計(CLABSI)'!$B$8:$M$8"
                    from openpyxl.chart.series import SeriesLabel
                    from openpyxl.chart.data_source import StrRef
                    s0.tx = SeriesLabel(strRef=StrRef("'年間集計(CLABSI)'!$A$8"))
                    ch.series = [s0]
                    ch.y_axis.scaling.max = None     # 元のグラフの上限20が残っていて、20を超える月が切れていた
                    runs = [r for p in ch.title.tx.rich.p for r in (p.r or [])]
                    if len(runs) >= 3:
                        runs[0].t, runs[1].t, runs[2].t = "CLABSI", "発生密度率", "(1000device-days)"
                    done.append("年間集計(CLABSI)：#REF! で何も表示されていなかったグラフを発生密度率（8行目）に修正")
    if "上半期 (CLABSI)" in wb.sheetnames:
        ws = wb["上半期 (CLABSI)"]
        if "発生件数" in str(ws["A7"].value or ""):
            for i in range(6):
                c = gl(i + 2)
                ws[f"{c}7"] = f"='年間集計(CLABSI)'!{c}7"
            done.append("上半期 (CLABSI)：件数（4〜6月が0の固定値、7〜9月が空欄）を年間集計から参照")
    return done
