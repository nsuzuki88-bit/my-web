# -*- coding: utf-8 -*-
"""CLABSI判定シート

  セクションA：血流感染（BSI）の判定 … 血液培養が陽性になった（またはCSEPを疑う）患者だけを
              1行ずつ入力する。患者IDからAllのブロックを探し、DOE・ICU在室日・中心ライン留置日数・
              RIT を自動で求めて JHAIS（Ver.2.5）の基準で判定する。
  セクションB：中心ライン（C有）の患者一覧 … Allシートで一度でも「C有」と入力された患者を
              CL初日の早い順に自動で集める（Allは日付順に入力するので、新しい患者は末尾に付く）。

判定の入力を患者一覧の行に直接持たせないのは、一覧が数式で並ぶため、Allの修正で並びが
変わると入力と患者がずれるから。BSIの入力は患者IDで患者を指定するので、並びに左右されない。
"""
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter as gl
from common import *

SHEET = "CLABSI判定"
BACT = ["認定病原体", "皮膚常在菌（2セット以上）", "皮膚常在菌（1セットのみ）",
        "陰性・未実施", "除外病原体"]
CRIT = ("LCBI-1", "LCBI-2", "CSEP")          # 一次血流感染として数える基準

# セクションA（BSIイベント）の列
EV = [  # (列見出し, 幅, 入力か)
    ("No", 5, False), ("患者ID\n(入力)", 12, True), ("血液培養\n採取日(入力)", 12, True),
    ("検出菌名\n(入力)", 15, True), ("菌の分類\n(選択)", 21, True),
    ("症状の初日\n(入力)", 11, True), ("2次性BSI\n(選択)", 11, True),
    ("CSEP：医師が\n敗血症治療を開始", 13, True), ("ブロック", 9, False), ("氏名", 13, False),
    ("基準", 22, False), ("DOE", 11, False), ("ICU在室\n(DOE時点)", 9, False),
    ("CL留置日数\n(DOEか前日)", 10, False), ("判定", 30, False), ("備考(入力)", 26, True),
]
# セクションB（中心ライン留置患者の一覧）の列
LS = ["No", "患者ID", "氏名", "ブロック", "部屋\n(CL初日)", "ICU入室日", "CL初日", "CL最終日",
      "CL日数", "CLABSI対象\n開始日(CL3日目)", "血流感染の判定"]

# 非表示の作業列
S_K, S_ID, S_NM, S_CLN, S_CL1, S_CLZ, S_ICU1, S_ICUZ, S_KEY, S_CL3 = range(CL_SCAN_C0, CL_SCAN_C0 + 10)
H_P, H_ROW, H_ICU0, H_ICU1, H_RICU0, H_RICU1, H_CL0, H_RCL0, H_RCL1, H_RIT, H_NEW = \
    range(CL_EV_H0, CL_EV_H0 + 11)


def build(wb):
    ws = wb.create_sheet(SHEET)
    ws.sheet_properties.tabColor = "7030A0"
    ws.sheet_view.showGridLines = False
    dates = f"All!$C${DATE_ROW}:${gl(CN)}${DATE_ROW}"
    sc = lambda col: f"${gl(col)}${CL_B_TOP}:${gl(col)}${CL_B_BOT}"      # 作業列（ブロックごと）
    ev = lambda col: f"${gl(col)}${CL_EV_TOP}:${gl(col)}${CL_EV_BOT}"    # セクションAの列

    ws["A1"] = "CLABSI 判定（Allシートの「C有」から中心ライン留置患者を自動で収集）"
    ws["A1"].font = F_TITLE
    ws["A1"].fill = fill("7030A0")
    ws.merge_cells("A1:P1")
    ws.row_dimensions[1].height = 27.75
    notes = [
        "① Allシートで「C有」と入力された患者は、下のセクションBに自動で集まります（CL初日の早い順）。"
        "② 血液培養が陽性になった患者（またはCSEPを疑う患者）だけを、セクションAに1行ずつ入力します（黄色のセル）。",
        "判定（JHAIS Ver.2.5）：LCBI-1／LCBI-2／CSEPを満たし、DOEか前日にCL3日目以降の中心ラインがあればCLABSI。"
        "入室2日以内のDOEはPOA・転棟元に帰属、DOEから14日間（RIT）の再陽性は新規にしません。",
    ]
    for i, text in enumerate(notes):
        c = ws.cell(2 + i, 1, text)
        c.font = F_NOTE
        c.alignment = A_LW
        ws.merge_cells(start_row=2 + i, start_column=1, end_row=2 + i, end_column=16)
        ws.row_dimensions[2 + i].height = 26
    c = ws.cell(4, 1, (f'="中心ライン（C有）の患者 "&COUNT({sc(S_KEY)})&"名 ／ 延べ中心ライン使用日数 "'
                       f'&SUM({sc(S_CLN)})&"日 ／ CLABSI "&COUNTIF({ev(15)},"CLABSI")'
                       f'&"件（月別の使用比・発生率は「CLABSI集計」シート）"'))
    c.font = Font(bold=True, size=10, color="7030A0")
    ws.merge_cells("A4:P4")

    # ============ セクションA：BSIイベントの判定 ==========================
    ws.cell(5, 1, "A. 血流感染（BSI）の判定　― 血液培養陽性（またはCSEPを疑う）患者だけを1行ずつ入力 ―")
    ws.cell(5, 1).font = F_SEC
    ws.merge_cells("A5:M5")
    c = ws.cell(5, 15, f'=HYPERLINK("#\'{SHEET}\'!A{CL_B_HDR}","▼ セクションB（患者一覧）へ")')
    c.font = Font(size=9, color="0563C1", underline="single")
    for i, (t, w, is_in) in enumerate(EV):
        style_cell(ws.cell(6, i + 1, t), "7F6000" if is_in else NAVY, F_HDR, A_CW)
        ws.column_dimensions[gl(i + 1)].width = w
    ws.row_dimensions[6].height = 42

    N_D = N_ALLD                                   # All!$C$1:$NC$last（日付列だけ）
    for e in range(CL_EV_TOP, CL_EV_BOT + 1):
        n = e - CL_EV_TOP + 1
        H = lambda col: f"${gl(col)}{e}"

        def rng(q):
            """このブロックのC有無行の、日付列1日目〜q日目"""
            return f"INDEX({N_D},{H(H_ROW)},1):INDEX({N_D},{H(H_ROW)},{q})"

        def at(q):
            return f"INDEX({N_D},{H(H_ROW)},{q})"

        def run(q, cond):
            """q日目で終わる連続日数（cond＝その日が条件を満たさない、の式）"""
            r = rng(q)
            return f"{q}-IFERROR(LOOKUP(2,1/({cond(r)}),COLUMN({r})-2),0)"

        not_icu = lambda r: f'({r}<>"C有")*({r}<>"C無")'
        not_cl = lambda r: f'({r}<>"C有")'
        P, P1 = H(H_P), f"({H(H_P)}-1)"

        fallback = (f'IFERROR(INDEX({sc(S_K)},MATCH($B{e}&"",INDEX({sc(S_ID)}&"",0),0)),"")')
        by_doe = (f'INDEX({sc(S_K)},MATCH(TRUE,INDEX(({sc(S_ID)}&""=$B{e}&"")'
                  f'*({sc(S_ICU1)}<=$L{e})*({sc(S_ICUZ)}+1>=$L{e})>0,0),0))')
        f = {
            1: n,
            # 同じ患者IDのブロックが複数ある（ICU再入室など）ときは、DOEがその在室期間
            # （入室日〜退室翌日）に入るブロックを選ぶ
            9: f'=IF($B{e}="","",IF($L{e}="",{fallback},IFERROR({by_doe},{fallback})))',
            10: f'=IF($I{e}="","",INDEX({sc(S_NM)},$I{e}))',
            11: (f'=IF(AND($C{e}="",$F{e}=""),"",'
                 f'IF($E{e}="{BACT[0]}",IF($C{e}="","（採取日を入力）","LCBI-1"),'
                 f'IF($E{e}="{BACT[1]}",IF($C{e}="","（採取日を入力）",'
                 f'IF($F{e}="","非該当（症状なし）",'
                 f'IF(ABS($F{e}-$C{e})<=3,"LCBI-2","非該当（症状がIWP外）"))),'
                 f'IF($E{e}="{BACT[2]}","非該当（汚染菌）",'
                 f'IF($E{e}="{BACT[3]}",IF($H{e}<>"はい","非該当（CSEP基準を満たさない）",'
                 f'IF($F{e}="","（症状の初日を入力）","CSEP")),'
                 f'IF($E{e}="{BACT[4]}","非該当（LCBI除外病原体）","（菌の分類を選択）"))))))'),
            12: (f'=IF($K{e}="LCBI-1",$C{e},IF($K{e}="LCBI-2",MIN($C{e},$F{e}),'
                 f'IF($K{e}="CSEP",$F{e},"")))'),
            13: f'=IF({P}="","",IF({H(H_ICU0)}=1,{H(H_RICU0)},IF({H(H_ICU1)}=1,{H(H_RICU1)}+1,0)))',
            14: f'=IF({P}="","",IF({H(H_CL0)}=1,{H(H_RCL0)},{H(H_RCL1)}))',
            15: (f'=IF(AND($B{e}="",$C{e}="",$F{e}=""),"",'
                 f'IF($B{e}="","⚠患者IDを入力",'
                 f'IF($I{e}="","⚠この患者IDはAllにありません",'
                 f'IF(NOT(OR($K{e}="LCBI-1",$K{e}="LCBI-2",$K{e}="CSEP")),$K{e},'
                 f'IF($G{e}="はい","2次性BSI（CLABSIではない）",'
                 f'IF({H(H_RIT)}=1,"RIT内（前回のBSIに菌を追加）",'
                 f'IF({P}="","⚠DOEが年度の範囲外",'
                 f'IF($M{e}=0,"ICU外（ICUのCLABSIではない）",'
                 f'IF($M{e}<=2,"入室2日以内（POA・転棟元に帰属）",'
                 f'IF($N{e}<3,"BSI（CL非関連）","CLABSI"))))))))))'),
            # ---- 非表示の作業列 ----
            H_P: f'=IF(OR($I{e}="",$L{e}=""),"",IFERROR(MATCH($L{e},{dates},0),""))',
            H_ROW: f'=IF($I{e}="","",{BLOCK0}+{BSTEP}*($I{e}-1)+3)',
            H_ICU0: f'=IF({P}="",0,IF(OR({at(P)}="C有",{at(P)}="C無"),1,0))',
            H_ICU1: (f'=IF(OR({P}="",N({P})<2),0,'
                     f'IF(OR({at(P1)}="C有",{at(P1)}="C無"),1,0))'),
            H_RICU0: f'=IF({H(H_ICU0)}=0,0,{run(P, not_icu)})',
            H_RICU1: f'=IF({H(H_ICU1)}=0,0,{run(P1, not_icu)})',
            H_CL0: f'=IF({P}="",0,IF({at(P)}="C有",1,0))',
            H_RCL0: f'=IF({H(H_CL0)}=0,0,{run(P, not_cl)})',
            H_RCL1: (f'=IF(OR({P}="",N({P})<2),0,IF({at(P1)}<>"C有",0,{run(P1, not_cl)}))'),
            # RIT：同じブロックで、上の行にDOEから13日以内の「新規の一次BSI」があるか
            H_RIT: (0 if e == CL_EV_TOP else
                    f'=IF(OR($I{e}="",$L{e}=""),0,IF(COUNTIFS($I${CL_EV_TOP}:$I{e-1},$I{e},'
                    f'$L${CL_EV_TOP}:$L{e-1},">="&($L{e}-13),$L${CL_EV_TOP}:$L{e-1},"<="&$L{e},'
                    f'${gl(H_NEW)}${CL_EV_TOP}:${gl(H_NEW)}{e-1},1)>0,1,0))'),
            H_NEW: (f'=IF(AND(OR($K{e}="LCBI-1",$K{e}="LCBI-2",$K{e}="CSEP"),$I{e}<>"",'
                    f'$G{e}<>"はい",{H(H_RIT)}=0),1,0)'),
        }
        fmts = {3: "yyyy/m/d", 6: "yyyy/m/d", 12: "yyyy/m/d",
                13: '0"日目";-0;"ICU外"', 14: '0"日";-0;"なし"'}
        for col, v in f.items():
            c = ws.cell(e, col, v)
            if col <= 16:
                is_in = EV[col - 1][2]
                style_cell(c, YELLOW if is_in else AUTO, F_BODY,
                           A_L if col in (4, 15, 16) else A_C, numfmt=fmts.get(col, "General"))
        for col in (2, 3, 4, 5, 6, 7, 8, 16):
            if ws.cell(e, col).value is None:
                style_cell(ws.cell(e, col), YELLOW, F_BODY,
                           A_L if col in (4, 16) else A_C, numfmt=fmts.get(col, "General"))
        style_cell(ws.cell(e, 1), BAND, Font(bold=True, size=9), A_C)
        style_cell(ws.cell(e, 15), GREEN, F_RESULT, A_L)

    ws.conditional_formatting.add(
        f"A{CL_EV_TOP}:P{CL_EV_BOT}",
        Rule(type="expression", formula=[f'$O{CL_EV_TOP}="CLABSI"'],
             dxf=DifferentialStyle(font=Font(bold=True), fill=cf_fill(PINK)), stopIfTrue=False))

    # 入力規則とヒント
    dv_date = DataValidation(
        type="date", operator="greaterThan", formula1="36526", allow_blank=True,
        errorStyle="warning", showErrorMessage=True, showInputMessage=True,
        errorTitle="日付の確認", error="日付（例：2026/9/18）で入力してください。",
        promptTitle="日付", prompt="例：2026/9/18")
    dv_bact = DataValidation(
        type="list", formula1='"' + ",".join(BACT) + '"', allow_blank=True, showInputMessage=True,
        promptTitle="菌の分類",
        prompt="認定病原体＝皮膚常在菌以外（黄色ブドウ球菌・腸内細菌目・緑膿菌・カンジダ等）。"
               "皮膚常在菌＝CoNS・コリネバクテリウム・バチルス・キューティバクテリウム・"
               "ミクロコッカス・緑色連鎖球菌等（別々の採血で同じ菌が2セット以上ならLCBI-2の候補）。")
    dv_sym = DataValidation(
        type="date", operator="greaterThan", formula1="36526", allow_blank=True,
        errorStyle="warning", showErrorMessage=True, showInputMessage=True,
        promptTitle="症状の初日",
        prompt="LCBI-2：発熱(>38.0℃)・悪寒戦慄・低血圧のいずれかが、採取日±3日（IWP）内で"
               "最初に見られた日。CSEP：発熱・低血圧(収縮期≦90mmHg)・乏尿(<20mL/時)の最初の日。")
    dv_sec = DataValidation(
        type="list", formula1='"はい,いいえ"', allow_blank=True, showInputMessage=True,
        promptTitle="2次性BSI",
        prompt="他部位の感染定義（JHAIS/NHSN）を満たし、2次血流感染帰属期間内に同じ菌が"
               "血液から出た（または血液培養がその定義の要素）なら「はい」。")
    dv_csep = DataValidation(
        type="list", formula1='"はい,いいえ"', allow_blank=True, showInputMessage=True,
        promptTitle="CSEP（臨床的敗血症）",
        prompt="血液培養が陰性・未実施で、他に原因も他部位の感染もなく、医師が敗血症として"
               "治療を開始したら「はい」。JHAISはCSEPを含めて報告します。")
    for dv in (dv_date, dv_bact, dv_sym, dv_sec, dv_csep):
        ws.add_data_validation(dv)
    dv_date.add(f"C{CL_EV_TOP}:C{CL_EV_BOT}")
    dv_bact.add(f"E{CL_EV_TOP}:E{CL_EV_BOT}")
    dv_sym.add(f"F{CL_EV_TOP}:F{CL_EV_BOT}")
    dv_sec.add(f"G{CL_EV_TOP}:G{CL_EV_BOT}")
    dv_csep.add(f"H{CL_EV_TOP}:H{CL_EV_BOT}")

    # ============ セクションB：中心ライン（C有）の患者一覧 ================
    ws.cell(CL_B_HDR, 1, f"B. 中心ライン（C有）の患者一覧　※Allシート先頭{NBLK}ブロックから、"
                         f"一度でも「C有」と入力された患者を自動で集めます（CL初日の早い順）")
    ws.cell(CL_B_HDR, 1).font = F_SEC
    ws.merge_cells(start_row=CL_B_HDR, start_column=1, end_row=CL_B_HDR, end_column=13)
    c = ws.cell(CL_B_HDR, 15, f'=HYPERLINK("#\'{SHEET}\'!A1","▲ セクションA（判定）へ")')
    c.font = Font(size=9, color="0563C1", underline="single")
    for i, t in enumerate(LS):
        style_cell(ws.cell(CL_B_HDR + 1, i + 1, t), NAVY, F_HDR, A_CW)
    ws.row_dimensions[CL_B_HDR + 1].height = 32

    for k in range(1, NBLK + 1):
        r = CL_B_TOP + k - 1
        cr = all_c(k)
        row = f"All!$C${cr}:${gl(CN)}${cr}"
        # ---- 非表示の作業列：ブロック k のスキャン ----
        scan = {
            S_K: k,
            S_ID: (f'=IF(INDEX(All!$A$1:$A${ALL_LAST_ROW},{all_id(k)})="","",'
                   f'INDEX(All!$A$1:$A${ALL_LAST_ROW},{all_id(k)}))'),
            S_NM: (f'=IF(INDEX(All!$A$1:$A${ALL_LAST_ROW},{all_name(k)})="","",'
                   f'INDEX(All!$A$1:$A${ALL_LAST_ROW},{all_name(k)}))'),
            S_CLN: f'=COUNTIF({row},"C有")',
            S_CL1: f'=IF(${gl(S_CLN)}{r}=0,"",MATCH("C有",{row},0))',
            S_CLZ: f'=IF(${gl(S_CLN)}{r}=0,"",LOOKUP(2,1/({row}="C有"),{dates}))',
            S_ICU1: (f'=IFERROR(INDEX({dates},1,MIN(IFERROR(MATCH("C有",{row},0),999),'
                     f'IFERROR(MATCH("C無",{row},0),999))),0)'),
            S_ICUZ: f'=IFERROR(LOOKUP(2,1/(({row}="C有")+({row}="C無")),{dates}),0)',
            S_KEY: f'=IF(${gl(S_CLN)}{r}=0,"",${gl(S_CL1)}{r}*10000+${gl(S_K)}{r})',
            # 中心ラインが3日連続した最初の日（＝CLABSIの判定対象になる最初の日）
            S_CL3: (f'=IF(${gl(S_CLN)}{r}<3,"",IFERROR(INDEX({dates},1,MATCH(TRUE,INDEX('
                    f'(All!$E${cr}:${gl(CN)}${cr}="C有")*(All!$D${cr}:${gl(CN-1)}${cr}="C有")'
                    f'*(All!$C${cr}:${gl(CN-2)}${cr}="C有")>0,0),0)+2),""))'),
        }
        for col, v in scan.items():
            ws.cell(r, col, v)

        # ---- 一覧（n番目の患者）----
        n = k
        blk = f"$D{r}"
        pick = lambda col: f'=IF({blk}="","",INDEX({sc(col)},{blk}))'
        room_row = f"{BLOCK0}+{BSTEP}*({blk}-1)+4"
        room_col = f"INDEX({sc(S_CL1)},{blk})+{C0-1}"
        vals = [
            (1, f'=IF({blk}="","",{n})', "0"),
            (2, pick(S_ID), "General"),
            (3, pick(S_NM), "General"),
            (4, f'=IFERROR(MOD(SMALL({sc(S_KEY)},{n}),10000),"")', "0"),
            (5, (f'=IF({blk}="","",IF(INDEX({N_ALL},{room_row},{room_col})="","",'
                 f'INDEX({N_ALL},{room_row},{room_col})))'), "General"),
            (6, pick(S_ICU1), "yyyy/m/d"),
            (7, f'=IF({blk}="","",INDEX({dates},1,INDEX({sc(S_CL1)},{blk})))', "yyyy/m/d"),
            (8, pick(S_CLZ), "yyyy/m/d"),
            (9, pick(S_CLN), "0"),
            (10, pick(S_CL3), "yyyy/m/d"),
            (11, (f'=IF({blk}="","",IF(COUNTIFS({ev(9)},{blk},{ev(15)},"CLABSI")>0,'
                  f'"CLABSI "&COUNTIFS({ev(9)},{blk},{ev(15)},"CLABSI")&"件",'
                  f'IF(COUNTIF({ev(9)},{blk})>0,"評価済（CLABSIなし）","")))'), "General"),
        ]
        for col, v, fmt in vals:
            style_cell(ws.cell(r, col, v), AUTO, F_BODY, A_C, numfmt=fmt)
        style_cell(ws.cell(r, 1), BAND, Font(bold=True, size=9), A_C)
        style_cell(ws.cell(r, 11), GREEN, F_RESULT, A_C)

    ws.conditional_formatting.add(
        f"A{CL_B_TOP}:K{CL_B_BOT}",
        Rule(type="expression", formula=[f'LEFT($K{CL_B_TOP},6)="CLABSI"'],
             dxf=DifferentialStyle(font=Font(bold=True), fill=cf_fill(PINK)), stopIfTrue=False))

    beyond = BLOCK0 + BSTEP * NBLK
    if beyond <= ALL_LAST_ROW:
        c = ws.cell(CL_B_BOT + 1, 1,
                    f'=IF(COUNTIF(All!$C${beyond}:${gl(CN)}${ALL_LAST_ROW},"C有")=0,'
                    f'"（スキャン範囲より後ろに「C有」はありません）",'
                    f'"⚠ ブロック{NBLK}より後ろにも「C有」があります。tools/build.py の第4引数で'
                    f'スキャン範囲を広げてください")')
        c.font = Font(size=9, bold=True, color=RED)
        ws.merge_cells(start_row=CL_B_BOT + 1, start_column=1,
                       end_row=CL_B_BOT + 1, end_column=11)

    for col in list(range(CL_SCAN_C0, CL_SCAN_C0 + 10)) + list(range(CL_EV_H0, CL_EV_H0 + 11)):
        ws.column_dimensions[gl(col)].hidden = True
    # 2つのセクションで見出しが違うので、行の固定はしない（上の「▼▲」リンクで移動する）
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:P{CL_EV_BOT}"
    return ws
