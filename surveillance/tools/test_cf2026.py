# -*- coding: utf-8 -*-
"""実データでAllシートの条件付き書式を検証する（式を同座標の別シートで評価）"""
import sys, openpyxl
from openpyxl.utils import get_column_letter as gl
import common, build
wb0 = openpyxl.load_workbook('new2026.xlsm', keep_vba=True, read_only=True)
g=build.probe(wb0['All']); g['NBLK']=min(g['NBLK'], common.NBLK_MAX)
common.configure(**g); common.configure(NCLIN=40, NSHEET=10)
from common import *
import sheet_all, ref_vae as R

BLOCKS = [414, 179, 248, 178, 323, 1, 2]     # DOEあり1件＋MV日数の多い例＋先頭

def make(src, dst):
    wb = openpyxl.load_workbook(src, keep_vba=True)
    for name, tmpl in (("_CF_DOE", sheet_all.doe_formula()),
                       ("_CF_VAC", sheet_all.vac_formula())):
        ws = wb.create_sheet(name)
        for k in BLOCKS:
            for off in (1, 2):
                r = all_v(k) + off
                for c in range(C0, CN + 1):
                    f = (tmpl.replace(f"COUNT(C{BLOCK0})", f"COUNT(All!{gl(c)}{r})")
                             .replace(f"C${DATE_ROW}", f"All!{gl(c)}${DATE_ROW}"))
                    ws.cell(r, c, "=" + f)
    wb.save(dst); print("wrote", dst)

def verify(path, src):
    dates, blocks = R.read_blocks(src)
    bymap = {b["k"]: b for b in blocks}
    wb = openpyxl.load_workbook(path, data_only=True)
    doe, vac = wb["_CF_DOE"], wb["_CF_VAC"]
    ok = ng = 0
    for k in BLOCKS:
        b = bymap[k]
        eps, tot = R.episodes(b)
        want_doe = set()
        want_vac = set()
        for (s, L) in eps:
            for d in R.does_for_episode(b, s, L):
                want_doe.add(dates[s + d - 1])
        # 全エピソードでVAC基準に合致する日（イベント期間の重複排除をしない版）
        for (s, L) in eps:
            n = min(L, MVROWS)
            F = [R.norm_fio2(b["f"][s + i]) for i in range(n)]
            P = [R.adj_peep(b["p"][s + i]) for i in range(n)]
            for i in range(n):
                if i + 1 < 3 or i + 1 >= n: continue
                f2, f1, f0, fp = F[i-2], F[i-1], F[i], F[i+1]
                p2, p1, p0, pp = P[i-2], P[i-1], P[i], P[i+1]
                if ((None not in (f2,f1,f0,fp) and f1 <= f2
                     and f0-f2 >= FIO2_RISE-1e-9 and fp-f2 >= FIO2_RISE-1e-9)
                    or (None not in (p2,p1,p0,pp) and p1 <= p2
                        and p0-p2 >= PEEP_RISE and pp-p2 >= PEEP_RISE)):
                    want_vac.add(dates[s + i])
        got_doe, got_cand = [], []
        for c in range(C0, CN + 1):
            r = all_f(k)
            dv, vv = doe.cell(r, c).value, vac.cell(r, c).value
            if doe.cell(all_p(k), c).value != dv or vac.cell(all_p(k), c).value != vv:
                ng += 1; print(f"  NG block{k} {gl(c)}: FiO2行とPEEP行で結果が違う")
            else: ok += 1
            if dv is True: got_doe.append(dates[c - C0])
            elif vv is True: got_cand.append(dates[c - C0])
        fmt = lambda xs: [x.strftime("%m/%d") for x in sorted(xs)]
        print(f"block{k:4d} ID={b['pid']}: 濃い赤={fmt(got_doe)}  薄い赤={fmt(got_cand)}")
        if set(got_doe) == want_doe: ok += 1
        else: ng += 1; print(f"  NG block{k} DOE: got {fmt(got_doe)} want {fmt(want_doe)}")
        if set(got_doe) | set(got_cand) == want_vac: ok += 1
        else: ng += 1; print(f"  NG block{k} VAC合致日: got {fmt(set(got_doe)|set(got_cand))} want {fmt(want_vac)}")
    print(f"\n==== ok={ok}  NG={ng} ====")
    return ng

if __name__ == "__main__":
    if sys.argv[1] == "make": make(sys.argv[2], sys.argv[3])
    else: sys.exit(1 if verify(sys.argv[2], sys.argv[3]) else 0)
