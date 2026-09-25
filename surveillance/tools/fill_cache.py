# -*- coding: utf-8 -*-
"""数式セルに計算結果（キャッシュ値）を書き込む

openpyxl で保存したブックは、数式セルの値が空（<v />）になる。Excel は開いたときに
全再計算する（fullCalcOnLoad）ので最終的には正しく表示されるが、
  ・保護ビュー（インターネットから取得したファイル）
  ・スマートフォンやメールのファイルプレビュー
では再計算されないため、一覧・臨床所見・判定シート・集計シートがすべて空欄に見える。

LibreOffice で全再計算したブックから値を読み、元のブックの <c> 要素に <v> を差し込む。
グラフも同じで、系列が参照するセルの値をグラフのキャッシュ（numCache／strCache）に書き込む
（openpyxl で作ったグラフはキャッシュが空、既存のグラフは元のファイルの古い値のままのため）。
数式・書式・VBA・グラフの設定など、値以外は一切変更しない。

  python3 fill_cache.py <build.pyの出力> <再計算済み(strip_charts後)> <出力>
"""
import html, re, sys, zipfile, datetime as dt
import openpyxl
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.datetime import to_excel

CELL = re.compile(r'<c r="([A-Z]+\d+)"([^>]*)>(<f[^>]*/>|<f[^>]*>[^<]*</f>)(<v\s*/>|<v>[^<]*</v>)?</c>')


def sheet_parts(z):
    """シート名 → ワークシートXMLのパス"""
    wbx = z.read("xl/workbook.xml").decode()
    relx = z.read("xl/_rels/workbook.xml.rels").decode()
    rels = {}
    for m in re.finditer(r"<Relationship [^>]*>", relx):
        t = m.group(0)
        rels[re.search(r'Id="([^"]+)"', t).group(1)] = re.search(r'Target="([^"]+)"', t).group(1)
    out = {}
    for m in re.finditer(r"<sheet [^>]*>", wbx):
        t = m.group(0)
        name = html.unescape(re.search(r'name="([^"]*)"', t).group(1))
        rid = re.search(r'r:id="([^"]+)"', t).group(1)
        out[name] = "xl/" + rels[rid].replace("/xl/", "").lstrip("/")
    return out


def cached(ws, coords):
    """再計算済みブックから、指定セルの (値, データ型) を取る"""
    rows = sorted({int(re.sub(r"[A-Z]+", "", c)) for c in coords})
    want = set(coords)
    got = {}
    for row in ws.iter_rows(min_row=rows[0], max_row=rows[-1]):
        for c in row:
            if getattr(c, "coordinate", None) in want:
                got[c.coordinate] = (c.value, c.data_type)
    return got


def v_xml(value, dtype):
    """<v> 要素と t 属性"""
    if dtype == "e":
        return ' t="e"', f"<v>{html.escape(str(value))}</v>"
    if isinstance(value, bool):
        return ' t="b"', f"<v>{int(value)}</v>"
    if isinstance(value, (dt.datetime, dt.date, dt.time, dt.timedelta)):
        return "", f"<v>{to_excel(value)!r}</v>"
    if isinstance(value, (int, float)):
        return "", f"<v>{value!r}</v>"
    if value is None:
        return ' t="str"', "<v></v>"
    return ' t="str"', f"<v>{html.escape(str(value), quote=False)}</v>"


# グラフの系列の参照（openpyxl は既定の名前空間で、Excel は c: を付けて書く。どちらにも対応）
REF = re.compile(r'<(?P<p>(?:c:)?)(?P<kind>numRef|strRef)><(?P=p)f>(?P<f>[^<]*)</(?P=p)f>'
                 r'(?P<cache><(?P=p)(?:numCache|strCache)>.*?</(?P=p)(?:numCache|strCache)>)?'
                 r'</(?P=p)(?P=kind)>', re.S)


def ref_values(wc, f):
    """'シート'!$B$5:$M$5 の値（行優先の1次元）。読めない参照は None"""
    f = html.unescape(f)
    if "!" not in f or "#REF" in f:
        return None
    sheet, rng = f.rsplit("!", 1)
    sheet = sheet.strip("'").replace("''", "'")
    if sheet not in wc.sheetnames:
        return None
    c1, r1, c2, r2 = range_boundaries(rng.replace("$", ""))
    rows = wc[sheet].iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2, values_only=True)
    return [v for row in rows for v in row]


def chart_cache(m, vals):
    p, kind = m.group("p"), m.group("kind")
    pts = []
    if kind == "numRef":
        fmt = re.search(r"<%sformatCode[^>]*>([^<]*)</%sformatCode>" % (p, p), m.group("cache") or "")
        for i, v in enumerate(vals):
            # 数値だけを入れる。#N/A・空欄・文字は点を置かない＝グラフに描かない
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                pts.append(f"<{p}pt idx=\"{i}\"><{p}v>{v!r}</{p}v></{p}pt>")
        head = (f"<{p}formatCode>{fmt.group(1) if fmt else 'General'}</{p}formatCode>"
                f"<{p}ptCount val=\"{len(vals)}\"/>")
        cache = f"<{p}numCache>{head}{''.join(pts)}</{p}numCache>"
    else:
        for i, v in enumerate(vals):
            if v is None or (isinstance(v, str) and v.startswith("#")):
                continue
            t = v.strftime("%Y/%m/%d") if isinstance(v, dt.datetime) else str(v)
            pts.append(f"<{p}pt idx=\"{i}\"><{p}v>{html.escape(t, quote=False)}</{p}v></{p}pt>")
        cache = f"<{p}strCache><{p}ptCount val=\"{len(vals)}\"/>{''.join(pts)}</{p}strCache>"
    return f"<{p}{kind}><{p}f>{m.group('f')}</{p}f>{cache}</{p}{kind}>"


def fill_charts(zs, wc):
    out, n = {}, 0
    for name in zs.namelist():
        if not re.match(r"xl/charts/chart\d+\.xml$", name):
            continue
        x = zs.read(name).decode("utf-8")

        def repl(m):
            nonlocal n
            vals = ref_values(wc, m.group("f"))
            if vals is None:
                return m.group(0)
            n += 1
            return chart_cache(m, vals)

        x2 = REF.sub(repl, x)
        if x2 != x:
            out[name] = x2.encode("utf-8")
    return out, n


def main(src, calc, dst):
    zs = zipfile.ZipFile(src)
    parts = sheet_parts(zs)
    wc = openpyxl.load_workbook(calc, data_only=True, read_only=True)
    new_xml, total = {}, 0
    for name, part in parts.items():
        x = zs.read(part).decode("utf-8")
        coords = [m.group(1) for m in CELL.finditer(x)]
        if not coords:
            continue
        vals = cached(wc[name], coords)

        def repl(m):
            coord, attrs, f = m.group(1), m.group(2), m.group(3)
            if coord not in vals:
                return m.group(0)
            t, v = v_xml(*vals[coord])
            attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
            return f'<c r="{coord}"{attrs}{t}>{f}{v}</c>'

        x2, n = CELL.subn(repl, x)
        new_xml[part] = x2.encode("utf-8")
        total += len(vals)
        print(f"  {name}: 数式 {len(coords)}セル / 値を書き込み {len(vals)}セル")

    charts, nref = fill_charts(zs, wc)
    new_xml.update(charts)
    print(f"  グラフ {len(charts)}個：系列の参照 {nref}か所のキャッシュを更新")

    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zd:
        for info in zs.infolist():
            data = new_xml.get(info.filename) or zs.read(info.filename)
            zd.writestr(info, data, compress_type=zipfile.ZIP_DEFLATED)
    print(f"wrote {dst}（計 {total}セル）")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3])
