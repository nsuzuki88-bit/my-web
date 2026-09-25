# -*- coding: utf-8 -*-
"""臨床所見に体温・WBCを入れたときの挙動を確認する

IVACは「①体温>38℃ or <36℃、あるいは WBC≥12,000 or ≤4,000/mm³」かつ
「②新規抗菌薬開始（4QAD以上）」の両方が必要（NHSN/JHAIS）。
体温・WBCだけではIVACにならず、VACのままであることを確かめる。
"""
import datetime as dt
import sys, openpyxl
from openpyxl.utils import get_column_letter as gl
from common import (LIST_A_TOP, NCLIN, clin_base, C0, CN, DATE_ROW,
                    all_id, all_name, all_v, all_f, all_p, all_rm, LIST_B_TOP)
import ref_vae as R
from make_tests import CASES

# 臨床所見の行オフセット
T_HI, T_LO, W_HI, W_LO, ABX, ABX_MEMO, PVAP, SPEC = range(8)

# (スロット, 説明, [(DOEからの日数ずれ, 行オフセット, 値), ...], 期待する判定)
SCEN = [
    (1, "体温38.6℃のみ（新規抗菌薬なし）",
     [(0, T_HI, 38.6)], "VAC"),
    (2, "WBC 15,200のみ（新規抗菌薬なし）",
     [(0, W_HI, 15200)], "VAC"),
    (3, "体温38.6℃ ＋ 新規抗菌薬",
     [(0, T_HI, 38.6), (0, ABX, "はい")], "IVAC"),
    (4, "WBC 15,200 ＋ 新規抗菌薬（ただしDOE+5日＝ウィンドウ外）",
     [(5, W_HI, 15200), (5, ABX, "はい")], "VAC"),
    (5, "体温38.0℃ちょうど ＋ 新規抗菌薬（>38℃なので異常ではない）",
     [(0, T_HI, 38.0), (0, ABX, "はい")], "VAC"),
    (6, "体温38.1℃ ＋ 新規抗菌薬",
     [(0, T_HI, 38.1), (0, ABX, "はい")], "IVAC"),
    (7, "WBC 12,000ちょうど ＋ 新規抗菌薬（≥12,000なので異常）",
     [(0, W_HI, 12000), (0, ABX, "はい")], "IVAC"),
    (8, "WBC 4,000ちょうど ＋ 新規抗菌薬（≤4,000なので異常）",
     [(0, W_LO, 4000), (0, ABX, "はい")], "IVAC"),
    (9, "最低体温35.9℃ ＋ 新規抗菌薬（<36℃なので異常）",
     [(0, T_LO, 35.9), (0, ABX, "はい")], "IVAC"),
]


def col_of(ws_all, date):
    for c in range(C0, CN + 1):
        v = ws_all.cell(DATE_ROW, c).value
        if isinstance(v, dt.datetime) and v.date() == date.date():
            return c
    raise ValueError(date)


def make(src, dst):
    # ① Allシートに教科書例を入れる（臨床所見はまだ入れない）
    wb = openpyxl.load_workbook(src)
    al, lst = wb["All"], wb["VAE対象一覧"]
    for k, desc, vmark, fio2, peep in CASES:
        al.cell(all_id(k), 1, f"T{k:02d}")
        al.cell(all_name(k), 1, f"検証{k:02d}")
        for d, (f, p) in enumerate(zip(fio2, peep)):
            c = C0 + d
            if f is None and p is None:
                continue
            if vmark:
                al.cell(all_v(k), c, vmark)
            al.cell(all_f(k), c, f)
            al.cell(all_p(k), c, p)
            al.cell(all_rm(k), c, "A-1")
        lst.cell(LIST_B_TOP + k - 1, 11, 70)
    wb.save(dst)

    # ② VACの並び順とDOEを求め、各スロットに所見を入れる
    _, ref = R.analyse_global(dst)
    vac = R.vac_patients(ref)
    wb = openpyxl.load_workbook(dst)
    al, cl = wb["All"], wb["臨床所見"]
    plan = []
    for slot, desc, entries, want in SCEN:
        if slot > len(vac):
            continue
        r = vac[slot - 1]
        doe = r["doe"][0]
        b = clin_base(slot)
        for shift, off, val in entries:
            c = col_of(al, doe) + shift
            cl.cell(b + off, c, val)
        plan.append((slot, r["k"], doe, desc, want))
    wb.save(dst)
    print("wrote", dst)
    for slot, k, doe, desc, want in plan:
        print(f"  スロット{slot} = ブロック{k} DOE {doe:%m/%d}: {desc} → 期待 {want}")
    return plan


def verify(path, plan):
    wb = openpyxl.load_workbook(path, data_only=True)
    lst = wb["VAE対象一覧"]
    ok = ng = 0
    print(f"\n{'スロット':<8}{'ブロック':<9}{'DOE':<8}{'臨床所見':<10}"
          f"{'判定':<7}{'期待':<7}  内容")
    print("-" * 118)
    for slot, k, doe, desc, want in plan:
        r = LIST_A_TOP + slot - 1
        got = lst.cell(r, 9).value
        entered = lst.cell(r, 12).value
        mark = "ok " if got == want else "NG "
        if got == want:
            ok += 1
        else:
            ng += 1
        print(f"{mark}{slot:<5}{k:<9}{doe:%m/%d}   {entered or '―':<10}"
              f"{got or '―':<7}{want:<7}  {desc}")
    print(f"\n==== ok={ok}  NG={ng} ====")
    return ng


if __name__ == "__main__":
    import json
    if sys.argv[1] == "make":
        plan = make(sys.argv[2], sys.argv[3])
        json.dump([[s, k, d.isoformat(), de, w] for s, k, d, de, w in plan],
                  open(sys.argv[4], "w"))
    else:
        raw = json.load(open(sys.argv[3]))
        plan = [(s, k, dt.datetime.fromisoformat(d), de, w) for s, k, d, de, w in raw]
        sys.exit(1 if verify(sys.argv[2], plan) else 0)
